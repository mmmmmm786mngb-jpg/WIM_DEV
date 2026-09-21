#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Diagnostika neodnoznachnyh UK i delty RSHB/Ivashkin."""

from __future__ import annotations

import pathlib
import re
import sys
from collections import Counter, defaultdict
from decimal import Decimal, InvalidOperation

from openpyxl import load_workbook

sys.stdout.reconfigure(encoding="utf-8", errors="replace")

BASE = pathlib.Path(r"c:\1c\Cursor_1c\WIM_DEV\bases\WIM_FIn\projects\IMDEV-7330 НовыеТесты")
FILES = {
    "9": BASE / "НДФЛ_Управление_24889_ПоНовому_Аванкор_9.xlsx",
    "10": BASE / "НДФЛ_Управление_24889_ПоНовому_Аванкор_10.xlsx",
    "11": BASE / "НДФЛ_Управление_24889_ПоНовому_Аванкор_11_РДУ.xlsx",
}
MONEY = ("СуммаДохода", "СуммаВычета", "НалогооблагаемаяCумма", "СуммаКУдержанию")
FIO_RE = re.compile(r"\(([^)]+)\)\s*$")
EPS = Decimal("0.01")


def safe_print(text: str) -> None:
    try:
        print(text)
    except UnicodeEncodeError:
        print(text.encode("ascii", "replace").decode("ascii"))


def to_dec(v) -> Decimal:
    if v is None or v == "":
        return Decimal("0")
    if isinstance(v, Decimal):
        return v
    if isinstance(v, int) and not isinstance(v, bool):
        return Decimal(v)
    if isinstance(v, float):
        return Decimal(str(round(v, 4)))
    s = str(v).strip().replace(" ", "").replace("\xa0", "").replace(",", ".")
    try:
        return Decimal(s)
    except InvalidOperation:
        return Decimal("0")


def parse_uk(path: pathlib.Path) -> dict:
    safe_print("parse " + path.name)
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    it = ws.iter_rows(values_only=True)
    header = [str(c) if c is not None else "" for c in next(it)]
    idx = {name: i for i, name in enumerate(header)}
    docs = {}
    for row in it:
        if not row:
            continue
        num = str(row[idx["НомерДокумента"]] or "").strip()
        if not num:
            continue
        port = str(row[idx.get("Портфель", 5)] or "").strip()
        d = docs.setdefault(num, {"ports": set(), "sums": {n: Decimal("0") for n in MONEY}})
        if port:
            d["ports"].add(port)
        for n in MONEY:
            i = idx.get(n)
            if i is not None:
                d["sums"][n] += to_dec(row[i] if i < len(row) else None)
    wb.close()
    for d in docs.values():
        names = []
        for p in d["ports"]:
            m = FIO_RE.search(p)
            if m:
                names.append(m.group(1).strip())
        d["client"] = Counter(names).most_common(1)[0][0] if names else ""
        d["empty"] = d["sums"]["СуммаДохода"] == 0 and d["sums"]["СуммаКУдержанию"] == 0
    return docs


def index_all(docs):
    idx = defaultdict(list)
    for num, d in docs.items():
        for p in d["ports"]:
            if num not in idx[p]:
                idx[p].append(num)
    return idx


def main() -> int:
    docs = {k: parse_uk(p) for k, p in FILES.items()}
    idx9 = index_all(docs["9"])
    collisions = sum(1 for p, nums in idx9.items() if len(nums) > 1)
    safe_print("A9 port names pointing to >1 UK doc: {0}".format(collisions))

    for lab in ("10", "11"):
        amb_nz = 0
        amb_empty = 0
        samples = []
        for num, d in docs[lab].items():
            votes = Counter()
            for p in d["ports"]:
                for n9 in idx9.get(p, []):
                    votes[n9] += 1
            if len(votes) > 1:
                top = votes.most_common(2)
                if top[0][1] == top[1][1]:
                    if d["empty"]:
                        amb_empty += 1
                    else:
                        amb_nz += 1
                        if len(samples) < 8:
                            samples.append((num, d, top))
        safe_print("A{0} ambiguous empty={1} nonempty={2}".format(lab, amb_empty, amb_nz))
        for num, d, top in samples:
            safe_print("  {0} {1} income={2} votes={3} ports={4}".format(
                num, d["client"], d["sums"]["СуммаДохода"], top, list(d["ports"])[:3]
            ))

    for label, num10, num9 in (("Ivashkin", "000000000025249", "000000000025249"),
                               ("RSHB", "000000000043543", "000000000047634")):
        d10 = docs["10"].get(num10)
        d9 = docs["9"].get(num9)
        safe_print("--- {0} ---".format(label))
        if d10:
            safe_print("A10 {0} {1} income={2} withhold={3} ports={4}".format(
                num10, d10["client"], d10["sums"]["СуммаДохода"],
                d10["sums"]["СуммаКУдержанию"], len(d10["ports"])
            ))
            safe_print("  ports: " + " | ".join(sorted(d10["ports"])[:8]))
        if d9:
            safe_print("A9  {0} {1} income={2} withhold={3} ports={4}".format(
                num9, d9["client"], d9["sums"]["СуммаДохода"],
                d9["sums"]["СуммаКУдержанию"], len(d9["ports"])
            ))
            safe_print("  ports: " + " | ".join(sorted(d9["ports"])[:8]))
        if d10 and d9:
            only10 = d10["ports"] - d9["ports"]
            only9 = d9["ports"] - d10["ports"]
            safe_print("  common={0} only10={1} only9={2}".format(
                len(d10["ports"] & d9["ports"]), len(only10), len(only9)
            ))
            if only10:
                safe_print("  only10: " + " | ".join(sorted(only10)[:6]))
            if only9:
                safe_print("  only9: " + " | ".join(sorted(only9)[:6]))

    for name, needle in (("Lvov", "Львов Павел"), ("Balandin", "Баландин Дмитрий")):
        safe_print("--- {0} ---".format(name))
        for lab in ("9", "10", "11"):
            hits = []
            for num, d in docs[lab].items():
                blob = " ".join(d["ports"]) + " " + d["client"]
                if needle in blob:
                    hits.append((num, d))
            safe_print("A{0} hits={1}".format(lab, len(hits)))
            for num, d in hits:
                safe_print("  {0} income={1} {2}".format(num, d["sums"]["СуммаДохода"], " | ".join(sorted(d["ports"]))))

    # Does 11 contain Ivashkin / RSHB?
    for needle in ("Ивашкин", "РСХБ"):
        for lab in ("9", "10", "11"):
            n = sum(1 for d in docs[lab].values() if needle in d["client"] or any(needle in p for p in d["ports"]))
            safe_print("{0} in A{1}: {2} UK docs".format(needle, lab, n))

    # Coverage of nonempty A9: union of A10/A11 ports
    ports10 = set()
    ports11 = set()
    for d in docs["10"].values():
        ports10 |= d["ports"]
    for d in docs["11"].values():
        ports11 |= d["ports"]
    miss_nz = []
    for num, d in docs["9"].items():
        if d["empty"]:
            continue
        if d["ports"] and d["ports"].isdisjoint(ports10 | ports11):
            miss_nz.append((num, d))
    safe_print("A9 nonempty UK with no port in 10/11: {0}".format(len(miss_nz)))
    for num, d in miss_nz[:8]:
        safe_print("  {0} {1} {2}".format(num, d["client"], list(d["ports"])[:2]))

    # Better match: exact port-set
    set9 = {}
    for num, d in docs["9"].items():
        key = frozenset(d["ports"])
        set9.setdefault(key, []).append(num)
    for lab in ("10", "11"):
        exact = 0
        miss = 0
        miss_nz = 0
        diffs = 0
        for num, d in docs[lab].items():
            hits = set9.get(frozenset(d["ports"]), [])
            if not hits:
                miss += 1
                if not d["empty"]:
                    miss_nz += 1
            else:
                exact += 1
                d9 = docs["9"][hits[0]]
                if any(abs(d["sums"][n] - d9["sums"][n]) >= EPS for n in MONEY):
                    diffs += 1
        safe_print("A{0} exact port-set match={1} miss={2} miss_nz={3} money_diff_on_exact={4}".format(
            lab, exact, miss, miss_nz, diffs
        ))

    return 0


if __name__ == "__main__":
    raise SystemExit(main())
