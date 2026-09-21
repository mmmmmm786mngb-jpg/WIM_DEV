#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Sverka Avancor 9 (polnyi), 10 (pustoi pul = ne-RDU), 11 (Roznichnoe DU).
Konsol - ASCII. HTML/TXT - UTF-8.
"""

from __future__ import annotations

import html
import pathlib
import re
import zipfile
from collections import Counter, defaultdict
from decimal import Decimal, InvalidOperation
from xml.etree import ElementTree as ET

from openpyxl import load_workbook

NS = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}

BASE = pathlib.Path(r"c:\1c\Cursor_1c\WIM_DEV\bases\WIM_FIn\projects\IMDEV-7330 НовыеТесты")
PROJ = pathlib.Path(
    r"c:\1c\Cursor_1c\WIM_DEV\bases\WIM_FIn\projects"
    r"\IMDEV-9458 Доработка обработки Формирование начислений НДФЛ"
)
OUT_TXT = PROJ / "Тестирование" / "reports" / "avancor9_10_11_report.txt"
OUT_HTML = PROJ / "Документация" / "imdev9458_avancor9_10_11_report.html"

WATCH = (
    "Львов Павел Глебович",
    "Баландин Дмитрий Викторович",
)
WATCH_SHORT = ("Львов Павел", "Баландин Дмитрий")

RUNS = {
    "9": {
        "uk": BASE / "НДФЛ_Управление_24889_ПоНовому_Аванкор_9.xlsx",
        "pf": BASE / "НДФЛ_Портфели_24889_ПоНовому_Аванкор_9.xlsx",
        "log": BASE / "Лог формирования начисдений НДФЛ_24889_ПоНовому_Аванкор_9.xlsx",
        "msg": BASE / "Сообщения_24889_поНовому_Аванкор_9.txt",
    },
    "10": {
        "uk": BASE / "НДФЛ_Управление_24889_ПоНовому_Аванкор_10.xlsx",
        "pf": BASE / "НДФЛ_Портфели_24889_ПоНовому_Аванкор_10.xlsx",
        "log": BASE / "Лог формирования начисдений НДФЛ_24889_ПоНовому_Аванкор_10.xlsx",
        "msg": BASE / "Сообщения_24889_поНовому_Аванкор_10.txt",
    },
    "11": {
        "uk": BASE / "НДФЛ_Управление_24889_ПоНовому_Аванкор_11_РДУ.xlsx",
        "pf": BASE / "НДФЛ_Портфели_24889_ПоНовому_Аванкор_11_РДУ.xlsx",
        "log": BASE / "Лог формирования начисдений НДФЛ_24889_ПоНовому_Аванкор_11_рду.xlsx",
        "msg": BASE / "Сообщения_24889_поНовому_Аванкор_11_РДУ.txt",
    },
}

MONEY_UK = (
    "СуммаДохода",
    "СуммаВычета",
    "НалогооблагаемаяCумма",
    "СуммаКУдержанию",
    "СуммаРанееУдержанногоНДФЛ",
)
MONEY_PF = (
    "СуммаДохода",
    "СуммаВычета",
    "НалогооблагаемаяCумма",
    "СуммаКУдержанию",
    "СуммаИсчисленногоНалога",
    "СуммаРанееУдержанногоНДФЛ",
)
EPS = Decimal("0.01")
FIO_RE = re.compile(r"\(([^)]+)\)\s*$")
DU_RE = re.compile(r"ДУ\s+(\d+)", re.IGNORECASE)


def safe_print(text: str) -> None:
    try:
        print(text)
    except UnicodeEncodeError:
        print(text.encode("ascii", "replace").decode("ascii"))


def to_dec(v) -> Decimal:
    if v is None or v == "":
        return Decimal("0")
    if isinstance(v, Decimal):
        return v
    if isinstance(v, bool):
        return Decimal("0")
    if isinstance(v, int):
        return Decimal(v)
    if isinstance(v, float):
        return Decimal(str(round(v, 4)))
    s = str(v).strip().replace(" ", "").replace("\xa0", "").replace(",", ".")
    if not s:
        return Decimal("0")
    try:
        return Decimal(s)
    except InvalidOperation:
        return Decimal("0")


def money_zero(sums: dict[str, Decimal], extra_tax: bool) -> bool:
    if sums.get("СуммаДохода", Decimal("0")) != 0:
        return False
    if sums.get("СуммаКУдержанию", Decimal("0")) != 0:
        return False
    if extra_tax and sums.get("СуммаИсчисленногоНалога", Decimal("0")) != 0:
        return False
    return True


def cell(row, idx: dict, name: str):
    i = idx.get(name)
    if i is None or i >= len(row):
        return None
    return row[i]


def client_from_ports(ports) -> str:
    names = []
    for p in ports:
        m = FIO_RE.search(p)
        if m:
            names.append(m.group(1).strip())
    if not names:
        return ""
    return Counter(names).most_common(1)[0][0]


def parse_ndfl(path: pathlib.Path, money_cols: tuple[str, ...], extra_tax: bool) -> dict:
    safe_print("parse " + path.name)
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    it = ws.iter_rows(values_only=True)
    header = next(it)
    header = [str(c) if c is not None else "" for c in header]
    idx = {name: i for i, name in enumerate(header)}
    docs: dict[str, dict] = {}
    nrows = 0
    for row in it:
        if row is None or all(c is None or str(c).strip() == "" for c in row):
            continue
        nrows += 1
        num = str(cell(row, idx, "НомерДокумента") or "").strip()
        ref = str(cell(row, idx, "Ссылка") or "").strip()
        if not num and ref:
            m = re.search(r"(\d{9,}|[0-9]+-\d+)", ref)
            if m:
                num = m.group(1)
        if not num:
            continue
        port = str(cell(row, idx, "Портфель") or cell(row, idx, "Объект") or "").strip()
        tch = str(cell(row, idx, "ТабличнаяЧасть") or "").strip()
        d = docs.setdefault(
            num,
            {
                "ref": ref,
                "ports": set(),
                "tch": set(),
                "rows": 0,
                "sums": {n: Decimal("0") for n in money_cols},
            },
        )
        d["rows"] += 1
        if ref and not d["ref"]:
            d["ref"] = ref
        if port:
            d["ports"].add(port)
        if tch:
            d["tch"].add(tch)
        for n in money_cols:
            d["sums"][n] += to_dec(cell(row, idx, n))
    wb.close()
    for d in docs.values():
        d["client"] = client_from_ports(d["ports"])
        d["empty"] = money_zero(d["sums"], extra_tax)
    nonempty = {k: v for k, v in docs.items() if not v["empty"]}
    empty = {k: v for k, v in docs.items() if v["empty"]}
    totals = {n: sum((d["sums"][n] for d in docs.values()), Decimal("0")) for n in money_cols}
    clients = {d["client"] for d in docs.values() if d["client"]}
    return {
        "docs": docs,
        "nonempty": nonempty,
        "empty": empty,
        "totals": totals,
        "clients": clients,
        "rows": nrows,
    }


def parse_log(path: pathlib.Path) -> dict:
    with zipfile.ZipFile(path) as z:
        if "xl/sharedStrings.xml" not in z.namelist():
            strings = []
        else:
            root = ET.fromstring(z.read("xl/sharedStrings.xml"))
            strings = [
                "".join(t.text or "" for t in si.findall(".//m:t", NS))
                for si in root.findall("m:si", NS)
            ]
    header = strings[0] if strings else ""
    err = 0
    m = re.search(r"Ошибок:\s*(\d+)", header)
    if m:
        err = int(m.group(1))
    return {
        "header": header.replace("\n", " | "),
        "errors": err,
        "uk_docs": sum(1 for s in strings if "по управляющей компании" in s),
        "pf_docs": sum(1 for s in strings if "Начисление НДФЛ по портфелю" in s),
    }


def parse_msg(path: pathlib.Path) -> dict:
    text = path.read_text(encoding="utf-8", errors="replace")
    lines = [ln for ln in text.splitlines() if ln.strip()]
    kinds = defaultdict(int)
    watch = {name: 0 for name in WATCH}
    for ln in lines:
        for name in WATCH:
            if name in ln:
                watch[name] += 1
        if "Невозможно распределить НДФЛ" in ln:
            kinds["distribute"] += 1
        elif "граница актуальности" in ln:
            kinds["border"] += 1
        elif "не проведен" in ln:
            kinds["unposted"] += 1
        elif "отсутствуют данные для расчета" in ln:
            kinds["empty_skip"] += 1
        else:
            kinds["other"] += 1
    return {"lines": len(lines), "kinds": dict(kinds), "watch": watch}


def fmt_n(n) -> str:
    if isinstance(n, Decimal):
        s = f"{n:,.2f}"
    else:
        s = f"{int(n):,}"
    return s.replace(",", " ")


def index_ports(dump: dict) -> dict[str, str]:
    idx: dict[str, str] = {}
    collisions = 0
    for num, d in dump["docs"].items():
        for p in d["ports"]:
            prev = idx.get(p)
            if prev and prev != num:
                collisions += 1
            else:
                idx[p] = num
    dump["port_collisions"] = collisions
    return idx


def match_by_ports(src: dict, dst: dict, dst_idx: dict[str, str]) -> dict:
    mapping = {}
    unmatched = []
    ambiguous = []
    for num, d in src["docs"].items():
        votes = Counter()
        unknown_ports = []
        for p in d["ports"]:
            hit = dst_idx.get(p)
            if hit:
                votes[hit] += 1
            else:
                unknown_ports.append(p)
        if not votes:
            unmatched.append(num)
            continue
        best, cnt = votes.most_common(1)[0]
        if len(votes) > 1 and votes.most_common(2)[1][1] == cnt:
            ambiguous.append(num)
            continue
        mapping[num] = {
            "dst": best,
            "votes": cnt,
            "unknown_ports": unknown_ports,
        }
    return {
        "map": mapping,
        "unmatched": unmatched,
        "ambiguous": ambiguous,
        "unmatched_nz": [n for n in unmatched if n in src["nonempty"]],
    }


def money_diffs(src: dict, dst: dict, mapping: dict, money_cols: tuple[str, ...]) -> list:
    diffs = []
    for snum, info in mapping.items():
        da = src["docs"][snum]
        db = dst["docs"][info["dst"]]
        delta = {}
        for n in money_cols:
            dlt = da["sums"][n] - db["sums"][n]
            if abs(dlt) >= EPS:
                delta[n] = dlt
        if delta:
            diffs.append(
                {
                    "src": snum,
                    "dst": info["dst"],
                    "client": da["client"] or db["client"],
                    "delta": delta,
                    "src_income": da["sums"].get("СуммаДохода", Decimal("0")),
                    "dst_income": db["sums"].get("СуммаДохода", Decimal("0")),
                }
            )
    return diffs


def compare_by_number(full: dict, subset: dict, money_cols: tuple[str, ...]) -> dict:
    ka, kb = set(full["docs"]), set(subset["docs"])
    only_full = sorted(ka - kb)
    only_sub = sorted(kb - ka)
    both = sorted(ka & kb)
    diffs = []
    for num in both:
        da, db = full["docs"][num], subset["docs"][num]
        delta = {}
        for n in money_cols:
            dlt = db["sums"][n] - da["sums"][n]
            if abs(dlt) >= EPS:
                delta[n] = dlt
        if delta:
            diffs.append({"num": num, "client": db["client"] or da["client"], "delta": delta})
    return {
        "only_full": only_full,
        "only_sub": only_sub,
        "both": both,
        "diffs": diffs,
        "only_full_nz": [n for n in only_full if n in full["nonempty"]],
        "only_sub_nz": [n for n in only_sub if n in subset["nonempty"]],
    }


def watch_docs(dump: dict) -> dict[str, list[dict]]:
    out = {name: [] for name in WATCH}
    for num, d in dump["docs"].items():
        blob = " ".join(d["ports"]) + " " + d["client"]
        for name, short in zip(WATCH, WATCH_SHORT):
            if name in blob or short in blob:
                out[name].append(
                    {
                        "num": num,
                        "client": d["client"],
                        "income": d["sums"].get("СуммаДохода", Decimal("0")),
                        "withhold": d["sums"].get("СуммаКУдержанию", Decimal("0")),
                        "ports": sorted(d["ports"]),
                        "empty": d["empty"],
                    }
                )
    return out


def load_run(label: str) -> dict:
    p = RUNS[label]
    safe_print("=== Avancor " + label + " ===")
    uk = parse_ndfl(p["uk"], MONEY_UK, extra_tax=False)
    pf = parse_ndfl(p["pf"], MONEY_PF, extra_tax=True)
    uk["port_index"] = index_ports(uk)
    pf["port_index"] = index_ports(pf)
    return {
        "uk": uk,
        "pf": pf,
        "log": parse_log(p["log"]),
        "msg": parse_msg(p["msg"]),
        "watch_uk": watch_docs(uk),
        "watch_pf": watch_docs(pf),
    }


def coverage(full: dict, parts: list[dict]) -> dict:
    covered = set()
    for part in parts:
        covered |= set(part["port_index"])
    missing = sorted(set(full["port_index"]) - covered)
    extra = []
    for part in parts:
        extra.extend(sorted(set(part["port_index"]) - set(full["port_index"])))
    docs_missing = []
    for num, d in full["docs"].items():
        if d["ports"] and d["ports"].isdisjoint(covered):
            docs_missing.append(num)
    return {
        "missing_ports": missing,
        "extra_ports": extra,
        "docs_missing": docs_missing,
        "docs_missing_nz": [n for n in docs_missing if n in full["nonempty"]],
        "covered_ports": len(covered),
        "full_ports": len(full["port_index"]),
    }


def overlap_docs(a: dict, b: dict) -> dict:
    ports_a = set(a["port_index"])
    ports_b = set(b["port_index"])
    common_ports = ports_a & ports_b
    only_a = ports_a - ports_b
    only_b = ports_b - ports_a
    docs_a = set()
    docs_b = set()
    for p in common_ports:
        docs_a.add(a["port_index"][p])
        docs_b.add(b["port_index"][p])
    return {
        "common_ports": len(common_ports),
        "only_a": len(only_a),
        "only_b": len(only_b),
        "docs_a": docs_a,
        "docs_b": docs_b,
    }


def build_txt(a9, a10, a11, m) -> str:
    lines = []
    a = lines.append
    a("=== AVANCOR 9 / 10 / 11 ===")
    a("A9 = polnyi nabor")
    a("A10 = pustoi pul (ne-RDU)")
    a("A11 = pul Roznichnoe DU")
    a("Match = po imenam portfelei, ne po nomeru UK-dokumenta")
    a("")
    a("--- Log / messages ---")
    for lab, run in (("9", a9), ("10", a10), ("11", a11)):
        a("A{0} log: {1} errors={2}".format(lab, run["log"]["header"][:180], run["log"]["errors"]))
        a("A{0} msg lines={1} kinds={2}".format(lab, run["msg"]["lines"], run["msg"]["kinds"]))
        a("A{0} msg watch={1}".format(lab, run["msg"]["watch"]))
    a("")
    a("--- UK counts ---")
    for lab, run in (("9", a9), ("10", a10), ("11", a11)):
        uk = run["uk"]
        a(
            "A{0} docs={1} empty={2} nonempty={3} clients={4} ports={5} income={6} withhold={7}".format(
                lab, len(uk["docs"]), len(uk["empty"]), len(uk["nonempty"]),
                len(uk["clients"]), len(uk["port_index"]),
                uk["totals"]["СуммаДохода"], uk["totals"]["СуммаКУдержанию"],
            )
        )
    a("")
    a("--- PF counts ---")
    for lab, run in (("9", a9), ("10", a10), ("11", a11)):
        pf = run["pf"]
        a(
            "A{0} docs={1} empty={2} nonempty={3} ports={4} income={5} withhold={6}".format(
                lab, len(pf["docs"]), len(pf["empty"]), len(pf["nonempty"]),
                len(pf["port_index"]),
                pf["totals"]["СуммаДохода"], pf["totals"]["СуммаКУдержанию"],
            )
        )
    a("")
    a("--- UK match to A9 by portfolio ---")
    for lab, mm in (("10", m["uk10"]), ("11", m["uk11"])):
        a(
            "A{0}: mapped={1} unmatched={2} (nz={3}) ambiguous={4} money_diffs={5}".format(
                lab, len(mm["map"]), len(mm["unmatched"]), len(mm["unmatched_nz"]),
                len(mm["ambiguous"]), len(m["uk" + lab + "_diff"]),
            )
        )
        if mm["unmatched_nz"]:
            a("  nonempty unmatched (first 15):")
            src = a10["uk"] if lab == "10" else a11["uk"]
            for num in mm["unmatched_nz"][:15]:
                d = src["docs"][num]
                a("    {0} {1} income={2} ports={3}".format(
                    num, d["client"], d["sums"]["СуммаДохода"], len(d["ports"])
                ))
        if m["uk" + lab + "_diff"]:
            a("  money diffs (first 15):")
            for item in m["uk" + lab + "_diff"][:15]:
                a("    src={0} a9={1} {2} {3}".format(
                    item["src"], item["dst"], item["client"], item["delta"]
                ))
    a("")
    a("--- PF match to A9 by portfolio ---")
    for lab, mm in (("10", m["pf10"]), ("11", m["pf11"])):
        a(
            "A{0}: mapped={1} unmatched={2} (nz={3}) money_diffs={4}".format(
                lab, len(mm["map"]), len(mm["unmatched"]), len(mm["unmatched_nz"]),
                len(m["pf" + lab + "_diff"]),
            )
        )
        if mm["unmatched_nz"][:10]:
            a("  nonempty unmatched (first 10): {0}".format(mm["unmatched_nz"][:10]))
        if m["pf" + lab + "_diff"][:10]:
            a("  money diffs (first 10):")
            for item in m["pf" + lab + "_diff"][:10]:
                a("    src={0} a9={1} {2} {3}".format(
                    item["src"], item["dst"], item["client"], item["delta"]
                ))
    a("")
    a("--- Coverage of A9 ports by A10+A11 ---")
    a("UK covered_ports={0}/{1} missing_ports={2} extra_ports={3} docs_without_any={4} (nz={5})".format(
        m["cov_uk"]["covered_ports"], m["cov_uk"]["full_ports"],
        len(m["cov_uk"]["missing_ports"]), len(m["cov_uk"]["extra_ports"]),
        len(m["cov_uk"]["docs_missing"]), len(m["cov_uk"]["docs_missing_nz"]),
    ))
    a("PF covered_ports={0}/{1} missing_ports={2} extra_ports={3} docs_without_any={4} (nz={5})".format(
        m["cov_pf"]["covered_ports"], m["cov_pf"]["full_ports"],
        len(m["cov_pf"]["missing_ports"]), len(m["cov_pf"]["extra_ports"]),
        len(m["cov_pf"]["docs_missing"]), len(m["cov_pf"]["docs_missing_nz"]),
    ))
    if m["cov_uk"]["docs_missing_nz"][:10]:
        a("  A9 UK nonempty not in 10/11 (first 10):")
        for num in m["cov_uk"]["docs_missing_nz"][:10]:
            d = a9["uk"]["docs"][num]
            a("    {0} {1} income={2}".format(num, d["client"], d["sums"]["СуммаДохода"]))
    if m["cov_pf"]["docs_missing_nz"][:10]:
        a("  A9 PF nonempty not in 10/11 (first 10):")
        for num in m["cov_pf"]["docs_missing_nz"][:10]:
            d = a9["pf"]["docs"][num]
            a("    {0} {1} income={2} ports={3}".format(
                num, d["client"], d["sums"]["СуммаДохода"], list(d["ports"])[:2]
            ))
    a("")
    a("--- Overlap A10 vs A11 (smeshannye pary) ---")
    a("UK common_ports={0} only10={1} only11={2} uk_docs10={3} uk_docs11={4}".format(
        m["ov_uk"]["common_ports"], m["ov_uk"]["only_a"], m["ov_uk"]["only_b"],
        len(m["ov_uk"]["docs_a"]), len(m["ov_uk"]["docs_b"]),
    ))
    a("PF common_ports={0} only10={1} only11={2} pf_docs10={3} pf_docs11={4}".format(
        m["ov_pf"]["common_ports"], m["ov_pf"]["only_a"], m["ov_pf"]["only_b"],
        len(m["ov_pf"]["docs_a"]), len(m["ov_pf"]["docs_b"]),
    ))
    a("")
    a("--- Number overlap (spravochno, UK nomer smenyaetsya) ---")
    a("UK 10 vs 9 by number: both={0} only10={1} diffs={2}".format(
        len(m["num_uk10"]["both"]), len(m["num_uk10"]["only_sub"]), len(m["num_uk10"]["diffs"])
    ))
    a("UK 11 vs 9 by number: both={0} only11={1} diffs={2}".format(
        len(m["num_uk11"]["both"]), len(m["num_uk11"]["only_sub"]), len(m["num_uk11"]["diffs"])
    ))
    a("PF 10 vs 9 by number: both={0} only10={1} diffs={2}".format(
        len(m["num_pf10"]["both"]), len(m["num_pf10"]["only_sub"]), len(m["num_pf10"]["diffs"])
    ))
    a("PF 11 vs 9 by number: both={0} only11={1} diffs={2}".format(
        len(m["num_pf11"]["both"]), len(m["num_pf11"]["only_sub"]), len(m["num_pf11"]["diffs"])
    ))
    a("")
    a("--- Watch ---")
    for name in WATCH:
        a(name + ":")
        for lab, run in (("9", a9), ("10", a10), ("11", a11)):
            items = run["watch_uk"][name]
            a("  UK A{0}: {1}".format(lab, "NO" if not items else "YES n=" + str(len(items))))
            for d in items:
                a("    num={0} empty={1} income={2} withhold={3} ports={4}".format(
                    d["num"], d["empty"], d["income"], d["withhold"], len(d["ports"])
                ))
            pitems = run["watch_pf"][name]
            a("  PF A{0}: {1}".format(lab, "NO" if not pitems else "YES n=" + str(len(pitems))))
            for d in pitems:
                a("    num={0} empty={1} income={2} withhold={3} {4}".format(
                    d["num"], d["empty"], d["income"], d["withhold"],
                    (d["ports"][:1] or [""])[0][:80],
                ))
    return "\n".join(lines) + "\n"


def build_html(a9, a10, a11, m) -> str:
    def n(x):
        return html.escape(fmt_n(x))

    uk9, uk10, uk11 = a9["uk"], a10["uk"], a11["uk"]
    pf9, pf10, pf11 = a9["pf"], a10["pf"], a11["pf"]

    bad_10 = (
        len(m["uk10"]["unmatched_nz"]) != 0
        or len(m["uk10_diff"]) != 0
        or len(m["pf10"]["unmatched_nz"]) != 0
        or len(m["pf10_diff"]) != 0
    )
    bad_11 = (
        len(m["uk11"]["unmatched_nz"]) != 0
        or len(m["uk11_diff"]) != 0
        or len(m["pf11"]["unmatched_nz"]) != 0
        or len(m["pf11_diff"]) != 0
    )
    miss9 = len(m["cov_uk"]["docs_missing_nz"]) + len(m["cov_pf"]["docs_missing_nz"])
    mixed_uk = len(m["ov_uk"]["docs_a"])
    mixed_pf = m["ov_pf"]["common_ports"]

    if not bad_10 and not bad_11 and miss9 == 0:
        verdict = (
            "Прогон 10 и прогон 11 — подмножества полного прогона 9: "
            "по портфелям документы находятся, суммы совпали, покрытие полного набора полное."
        )
        vcls = "ok"
    elif not bad_10 and not bad_11:
        verdict = (
            "Суммы найденных документов совпали с прогоном 9, но часть живых документов 9 "
            "не попала ни в пустой пул, ни в Розничное ДУ."
        )
        vcls = "warn"
    else:
        verdict = (
            "Есть расхождения: живые документы 10/11 без пары в 9, либо суммы по тем же портфелям не совпали."
        )
        vcls = "err"

    def diff_rows(items):
        if not items:
            return "<tr><td colspan='4'>Расхождений сумм нет.</td></tr>\n"
        rows = ""
        for item in items[:20]:
            parts = ", ".join("{0}: {1}".format(k, fmt_n(v)) for k, v in item["delta"].items())
            rows += (
                "<tr><td>{0}</td><td>{1}</td><td>{2}</td><td>{3}</td></tr>\n".format(
                    html.escape(item["src"]), html.escape(item["dst"]),
                    html.escape(item["client"] or "—"), html.escape(parts),
                )
            )
        return rows

    def extra_rows(run, mm):
        nums = mm["unmatched_nz"]
        if not nums:
            return "<tr><td colspan='4'>Нет. Все живые документы нашлись в прогоне 9.</td></tr>\n"
        rows = ""
        for num in nums[:20]:
            d = run["docs"][num]
            rows += (
                "<tr><td>{0}</td><td>{1}</td><td class='num'>{2}</td><td class='num'>{3}</td></tr>\n"
                .format(
                    html.escape(d["client"] or "—"), html.escape(num),
                    n(d["sums"]["СуммаДохода"]), n(d["sums"]["СуммаКУдержанию"]),
                )
            )
        if len(nums) > 20:
            rows += "<tr><td colspan='4'>Показаны первые 20 из {0}.</td></tr>\n".format(n(len(nums)))
        return rows

    miss_rows = ""
    for num in m["cov_uk"]["docs_missing_nz"][:15]:
        d = uk9["docs"][num]
        miss_rows += (
            "<tr><td>УК</td><td>{0}</td><td>{1}</td><td class='num'>{2}</td></tr>\n"
            .format(html.escape(num), html.escape(d["client"] or "—"), n(d["sums"]["СуммаДохода"]))
        )
    for num in m["cov_pf"]["docs_missing_nz"][:15]:
        d = pf9["docs"][num]
        miss_rows += (
            "<tr><td>Портфель</td><td>{0}</td><td>{1}</td><td class='num'>{2}</td></tr>\n"
            .format(html.escape(num), html.escape(d["client"] or "—"), n(d["sums"]["СуммаДохода"]))
        )
    if not miss_rows:
        miss_rows = "<tr><td colspan='4'>Нет. Каждый живой документ прогона 9 встретился в 10 или в 11.</td></tr>\n"

    mixed_rows = ""
    shown = 0
    for num in sorted(m["ov_uk"]["docs_a"])[:12]:
        d = uk10["docs"][num]
        mixed_rows += (
            "<tr><td>{0}</td><td>{1}</td><td class='num'>{2}</td><td class='num'>{3}</td></tr>\n"
            .format(
                html.escape(num), html.escape(d["client"] or "—"),
                n(d["sums"]["СуммаДохода"]), n(len(d["ports"])),
            )
        )
        shown += 1
    if not mixed_rows:
        mixed_rows = "<tr><td colspan='4'>Пересечения портфелей УК между прогонами 10 и 11 нет.</td></tr>\n"
    elif len(m["ov_uk"]["docs_a"]) > shown:
        mixed_rows += (
            "<tr><td colspan='4'>Показаны первые {0} из {1} документов УК прогона 10, "
            "у которых портфель есть и в прогоне 11.</td></tr>\n"
            .format(n(shown), n(len(m["ov_uk"]["docs_a"])))
        )

    watch_rows = ""
    for name in WATCH:
        for lab, run in (("9", a9), ("10", a10), ("11", a11)):
            items = run["watch_uk"][name]
            if not items:
                watch_rows += (
                    "<tr><td><strong>{0}</strong></td><td>{1}</td><td colspan='3'>нет документа УК</td></tr>\n"
                    .format(html.escape(name), lab)
                )
                continue
            for d in items:
                watch_rows += (
                    "<tr><td><strong>{0}</strong></td><td>{1}</td><td>{2}</td>"
                    "<td class='num'>{3}</td><td class='num'>{4}</td></tr>\n"
                    .format(
                        html.escape(name), lab, html.escape(d["num"]),
                        n(d["income"]), n(d["withhold"]),
                    )
                )

    return f"""<!DOCTYPE html>
<html lang="ru">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>IMDEV-9458. Сверка прогонов 9, 10 и 11</title>
<style>
  :root {{ --ok:#28a745; --err:#dc3545; --warn:#b78100; --ink:#212529; --soft:#6c757d;
    --bg:#f5f6f8; --card:#fff; --line:#dee2e6; }}
  * {{ box-sizing:border-box; }}
  body {{ margin:0; font-family:"Segoe UI", Arial, Helvetica, sans-serif; font-size:15px;
    line-height:1.45; color:var(--ink); background:var(--bg); }}
  .wrap {{ max-width:980px; margin:0 auto; padding:24px 16px 48px; }}
  header {{ background:#1f4a3a; color:#f7f1e8; border-radius:16px; padding:22px 24px; margin-bottom:20px; }}
  header .kicker {{ font-size:11px; letter-spacing:.1em; text-transform:uppercase; opacity:.75; margin-bottom:6px; }}
  header h1 {{ margin:0 0 8px; font-size:24px; font-family:Georgia, "Times New Roman", serif; }}
  header p {{ margin:0; color:#dce8e2; }}
  section {{ background:var(--card); border:1px solid var(--line); border-radius:12px; padding:18px 20px; margin-bottom:16px; }}
  h2 {{ margin:0 0 12px; font-size:18px; }}
  h3 {{ margin:16px 0 8px; font-size:16px; }}
  p, ul {{ margin:0 0 10px; }}
  ul {{ padding-left:20px; }}
  table {{ width:100%; border-collapse:collapse; margin:10px 0 4px; font-size:14px; }}
  th, td {{ border:1px solid var(--line); padding:8px 10px; text-align:left; vertical-align:top; }}
  th {{ background:#eef2f0; }}
  .num {{ text-align:right; font-variant-numeric:tabular-nums; white-space:nowrap; }}
  .ok {{ color:var(--ok); font-weight:600; }}
  .warn {{ color:var(--warn); font-weight:600; }}
  .err {{ color:var(--err); font-weight:600; }}
  .box-green {{ background:#e8f6ec; border:1px solid #b7e0c2; border-radius:10px; padding:12px 14px; margin:10px 0; }}
  .box-blue {{ background:#e7f4fb; border:1px solid #b9d7ea; border-radius:10px; padding:12px 14px; margin:10px 0; }}
  .box-yellow {{ background:#fff8e1; border:1px solid #f0d78c; border-radius:10px; padding:12px 14px; margin:10px 0; }}
  .grid {{ display:grid; grid-template-columns:repeat(auto-fit,minmax(150px,1fr)); gap:10px; margin:10px 0; }}
  .stat {{ background:#f8faf9; border:1px solid var(--line); border-radius:10px; padding:12px; }}
  .stat .v {{ font-size:22px; font-weight:700; }}
  .stat .l {{ font-size:12px; color:var(--soft); margin-top:4px; }}
  .muted {{ color:var(--soft); font-size:13px; }}
</style>
</head>
<body>
<div class="wrap">
<header>
  <div class="kicker">IMDEV-9458 · тесты 9, 10 и 11</div>
  <h1>Сверка полного набора, пустого пула и Розничного ДУ</h1>
  <p>Прогон 9 — все клиенты. Прогон 10 — пустой пул (не-РДУ). Прогон 11 — пул «Розничное ДУ».</p>
</header>

<section>
  <h2>1. Введение</h2>
  <p>Проверяем три утверждения. Прогон 10 должен быть не-РДУ частью полного расчёта. Прогон 11 — РДУ-частью. Вместе они должны покрыть прогон 9. Смешанные пары (есть и РДУ, и не-РДУ) могут попасть в оба отбора, потому что заполнение держит пару клиент + УК целиком.</p>
  <div class="box-green"><strong>Итог.</strong> <span class="{vcls}">{html.escape(verdict)}</span></div>
</section>

<section>
  <h2>2. Технология</h2>
  <p>Сверялись выгрузки документов по УК и по портфелю, лог формирования и сообщения. Колонки «Клиент» в выгрузках нет: ФИО берётся из скобок в наименовании портфеля. Номер документа УК между прогонами не считается ключом — после удаления и повторного создания он другой. Ключ сверки — наименование портфеля. Пустой документ: доход = 0 и к удержанию = 0 (для портфеля ещё исчисленный налог = 0).</p>
  <div class="box-blue">
    <ul>
      <li>Прогон 9: {html.escape(a9["log"]["header"][:220])}</li>
      <li>Прогон 10: {html.escape(a10["log"]["header"][:220])}</li>
      <li>Прогон 11: {html.escape(a11["log"]["header"][:220])}</li>
    </ul>
  </div>
</section>

<section>
  <h2>3. Статистика</h2>
  <div class="grid">
    <div class="stat"><div class="v">{n(len(uk9["docs"]))}</div><div class="l">УК, прогон 9</div></div>
    <div class="stat"><div class="v">{n(len(uk10["docs"]))}</div><div class="l">УК, прогон 10</div></div>
    <div class="stat"><div class="v">{n(len(uk11["docs"]))}</div><div class="l">УК, прогон 11</div></div>
    <div class="stat"><div class="v">{n(mixed_uk)}</div><div class="l">УК 10, пересечение с 11</div></div>
  </div>
  <h3>Документы по УК</h3>
  <table>
    <tr><th>Метрика</th><th class="num">Прогон 9</th><th class="num">Прогон 10</th><th class="num">Прогон 11</th></tr>
    <tr><td>Документов</td><td class="num">{n(len(uk9["docs"]))}</td><td class="num">{n(len(uk10["docs"]))}</td><td class="num">{n(len(uk11["docs"]))}</td></tr>
    <tr><td>Пустых</td><td class="num">{n(len(uk9["empty"]))}</td><td class="num">{n(len(uk10["empty"]))}</td><td class="num">{n(len(uk11["empty"]))}</td></tr>
    <tr><td>С суммами</td><td class="num">{n(len(uk9["nonempty"]))}</td><td class="num">{n(len(uk10["nonempty"]))}</td><td class="num">{n(len(uk11["nonempty"]))}</td></tr>
    <tr><td>Клиентов (из портфеля)</td><td class="num">{n(len(uk9["clients"]))}</td><td class="num">{n(len(uk10["clients"]))}</td><td class="num">{n(len(uk11["clients"]))}</td></tr>
    <tr><td>Портфелей в ТЧ</td><td class="num">{n(len(uk9["port_index"]))}</td><td class="num">{n(len(uk10["port_index"]))}</td><td class="num">{n(len(uk11["port_index"]))}</td></tr>
    <tr><td>Сумма дохода</td><td class="num">{n(uk9["totals"]["СуммаДохода"])}</td><td class="num">{n(uk10["totals"]["СуммаДохода"])}</td><td class="num">{n(uk11["totals"]["СуммаДохода"])}</td></tr>
    <tr><td>Сумма к удержанию</td><td class="num">{n(uk9["totals"]["СуммаКУдержанию"])}</td><td class="num">{n(uk10["totals"]["СуммаКУдержанию"])}</td><td class="num">{n(uk11["totals"]["СуммаКУдержанию"])}</td></tr>
  </table>
  <h3>Документы по портфелю</h3>
  <table>
    <tr><th>Метрика</th><th class="num">Прогон 9</th><th class="num">Прогон 10</th><th class="num">Прогон 11</th></tr>
    <tr><td>Документов</td><td class="num">{n(len(pf9["docs"]))}</td><td class="num">{n(len(pf10["docs"]))}</td><td class="num">{n(len(pf11["docs"]))}</td></tr>
    <tr><td>Пустых</td><td class="num">{n(len(pf9["empty"]))}</td><td class="num">{n(len(pf10["empty"]))}</td><td class="num">{n(len(pf11["empty"]))}</td></tr>
    <tr><td>С суммами</td><td class="num">{n(len(pf9["nonempty"]))}</td><td class="num">{n(len(pf10["nonempty"]))}</td><td class="num">{n(len(pf11["nonempty"]))}</td></tr>
    <tr><td>Сумма дохода</td><td class="num">{n(pf9["totals"]["СуммаДохода"])}</td><td class="num">{n(pf10["totals"]["СуммаДохода"])}</td><td class="num">{n(pf11["totals"]["СуммаДохода"])}</td></tr>
    <tr><td>Сумма к удержанию</td><td class="num">{n(pf9["totals"]["СуммаКУдержанию"])}</td><td class="num">{n(pf10["totals"]["СуммаКУдержанию"])}</td><td class="num">{n(pf11["totals"]["СуммаКУдержанию"])}</td></tr>
  </table>
  <p class="muted">Ошибок в логе: прогон 9 — {n(a9["log"]["errors"])}, прогон 10 — {n(a10["log"]["errors"])}, прогон 11 — {n(a11["log"]["errors"])}. Сообщений о распределении на первый портфель: {n(a9["msg"]["kinds"].get("distribute", 0))} / {n(a10["msg"]["kinds"].get("distribute", 0))} / {n(a11["msg"]["kinds"].get("distribute", 0))}.</p>
</section>

<section>
  <h2>4. Детали сверки</h2>
  <h3>10 и 11 против 9 по портфелю</h3>
  <table>
    <tr><th>Срез</th><th class="num">Сопоставлено</th><th class="num">Живые без пары в 9</th><th class="num">Расхождений сумм</th></tr>
    <tr><td>УК, прогон 10</td><td class="num">{n(len(m["uk10"]["map"]))}</td><td class="num">{n(len(m["uk10"]["unmatched_nz"]))}</td><td class="num">{n(len(m["uk10_diff"]))}</td></tr>
    <tr><td>УК, прогон 11</td><td class="num">{n(len(m["uk11"]["map"]))}</td><td class="num">{n(len(m["uk11"]["unmatched_nz"]))}</td><td class="num">{n(len(m["uk11_diff"]))}</td></tr>
    <tr><td>Портфель, прогон 10</td><td class="num">{n(len(m["pf10"]["map"]))}</td><td class="num">{n(len(m["pf10"]["unmatched_nz"]))}</td><td class="num">{n(len(m["pf10_diff"]))}</td></tr>
    <tr><td>Портфель, прогон 11</td><td class="num">{n(len(m["pf11"]["map"]))}</td><td class="num">{n(len(m["pf11"]["unmatched_nz"]))}</td><td class="num">{n(len(m["pf11_diff"]))}</td></tr>
  </table>

  <h3>Покрытие прогона 9 прогонами 10 + 11</h3>
  <table>
    <tr><th>Срез</th><th class="num">Портфелей в 9</th><th class="num">Есть в 10 или 11</th><th class="num">Живых документов 9 без портфеля в 10/11</th></tr>
    <tr><td>УК</td><td class="num">{n(m["cov_uk"]["full_ports"])}</td><td class="num">{n(m["cov_uk"]["covered_ports"])}</td><td class="num">{n(len(m["cov_uk"]["docs_missing_nz"]))}</td></tr>
    <tr><td>Портфель</td><td class="num">{n(m["cov_pf"]["full_ports"])}</td><td class="num">{n(m["cov_pf"]["covered_ports"])}</td><td class="num">{n(len(m["cov_pf"]["docs_missing_nz"]))}</td></tr>
  </table>
  <table>
    <tr><th>Вид</th><th>Номер в 9</th><th>Клиент</th><th class="num">Доход</th></tr>
    {miss_rows}
  </table>

  <div class="box-yellow">
    <strong>Смешанные пары.</strong> Портфелей, которые есть и в прогоне 10, и в прогоне 11:
    УК {n(m["ov_uk"]["common_ports"])}, отдельно по портфелю {n(mixed_pf)}.
    Документов УК прогона 10 с таким пересечением: {n(len(m["ov_uk"]["docs_a"]))}.
    Это ожидаемо, если у клиента в одной УК есть и РДУ, и не-РДУ: отбор держит пару целиком.
  </div>
  <table>
    <tr><th>Номер УК в 10</th><th>Клиент</th><th class="num">Доход</th><th class="num">Портфелей</th></tr>
    {mixed_rows}
  </table>

  <h3>Живые документы 10 без портфеля в 9</h3>
  <table>
    <tr><th>Клиент</th><th>Номер</th><th class="num">Доход</th><th class="num">К удержанию</th></tr>
    {extra_rows(uk10, m["uk10"])}
  </table>
  <h3>Живые документы 11 без портфеля в 9</h3>
  <table>
    <tr><th>Клиент</th><th>Номер</th><th class="num">Доход</th><th class="num">К удержанию</th></tr>
    {extra_rows(uk11, m["uk11"])}
  </table>
  <h3>Расхождения сумм УК 10 против 9</h3>
  <table>
    <tr><th>Номер 10</th><th>Номер 9</th><th>Клиент</th><th>Дельта (10 − 9)</th></tr>
    {diff_rows(m["uk10_diff"])}
  </table>
  <h3>Расхождения сумм УК 11 против 9</h3>
  <table>
    <tr><th>Номер 11</th><th>Номер 9</th><th>Клиент</th><th>Дельта (11 − 9)</th></tr>
    {diff_rows(m["uk11_diff"])}
  </table>
  <h3>Расхождения сумм портфеля 10 / 11 против 9</h3>
  <table>
    <tr><th>Номер среза</th><th>Номер 9</th><th>Клиент</th><th>Дельта</th></tr>
    {diff_rows(m["pf10_diff"] + m["pf11_diff"])}
  </table>

  <h3>Львов и Баландин</h3>
  <table>
    <tr><th>Клиент</th><th>Прогон</th><th>Номер УК</th><th class="num">Доход</th><th class="num">К удержанию</th></tr>
    {watch_rows}
  </table>
</section>

<section>
  <h2>5. Возможности</h2>
  <ul>
    <li>Если смешанных пар много — это не ошибка отбора по пулу, а зерно «клиент + УК».</li>
    <li>Обработка проверки проведенных документов по пулу РДУ показывает чужие портфели в ТЧ таких пар.</li>
    <li>Дельту по сумме имеет смысл открыть в базе 24889 по номеру из прогона 9 и сравнить ТЧ.</li>
  </ul>
</section>

<section>
  <h2>6. Выводы</h2>
  <ol>
    <li>Прогон 10 короче полного: остались клиенты с пустым пулом.</li>
    <li>Прогон 11 — основная масса полного набора, пул «Розничное ДУ».</li>
    <li>По сопоставленным портфелям суммы {"совпали" if not bad_10 and not bad_11 else "разошлись"}.</li>
    <li>Живых документов 9 без портфеля ни в 10, ни в 11: {n(miss9)}.</li>
    <li>Ошибок в логах трёх прогонов: {n(a9["log"]["errors"] + a10["log"]["errors"] + a11["log"]["errors"])}.</li>
  </ol>
</section>
<p class="muted">Отчёт по выгрузкам прогонов 9, 10 и 11, задача IMDEV-9458.</p>
</div>
</body>
</html>
"""


def main() -> int:
    a9 = load_run("9")
    a10 = load_run("10")
    a11 = load_run("11")

    uk10 = match_by_ports(a10["uk"], a9["uk"], a9["uk"]["port_index"])
    uk11 = match_by_ports(a11["uk"], a9["uk"], a9["uk"]["port_index"])
    pf10 = match_by_ports(a10["pf"], a9["pf"], a9["pf"]["port_index"])
    pf11 = match_by_ports(a11["pf"], a9["pf"], a9["pf"]["port_index"])

    m = {
        "uk10": uk10,
        "uk11": uk11,
        "pf10": pf10,
        "pf11": pf11,
        "uk10_diff": money_diffs(a10["uk"], a9["uk"], uk10["map"], MONEY_UK),
        "uk11_diff": money_diffs(a11["uk"], a9["uk"], uk11["map"], MONEY_UK),
        "pf10_diff": money_diffs(a10["pf"], a9["pf"], pf10["map"], MONEY_PF),
        "pf11_diff": money_diffs(a11["pf"], a9["pf"], pf11["map"], MONEY_PF),
        "cov_uk": coverage(a9["uk"], [a10["uk"], a11["uk"]]),
        "cov_pf": coverage(a9["pf"], [a10["pf"], a11["pf"]]),
        "ov_uk": overlap_docs(a10["uk"], a11["uk"]),
        "ov_pf": overlap_docs(a10["pf"], a11["pf"]),
        "num_uk10": compare_by_number(a9["uk"], a10["uk"], MONEY_UK),
        "num_uk11": compare_by_number(a9["uk"], a11["uk"], MONEY_UK),
        "num_pf10": compare_by_number(a9["pf"], a10["pf"], MONEY_PF),
        "num_pf11": compare_by_number(a9["pf"], a11["pf"], MONEY_PF),
    }

    txt = build_txt(a9, a10, a11, m)
    OUT_TXT.parent.mkdir(parents=True, exist_ok=True)
    OUT_TXT.write_text(txt, encoding="utf-8")
    OUT_HTML.write_text(build_html(a9, a10, a11, m), encoding="utf-8")
    safe_print(txt)
    safe_print("TXT OK")
    safe_print("HTML OK")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
