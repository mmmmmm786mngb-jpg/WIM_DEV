#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Sverka Avancor 10 (pustoi pul = ne-RDU) s Avancor 9 (polnyi nabor).
Konsol - ASCII. HTML/TXT - UTF-8.
"""

from __future__ import annotations

import html
import pathlib
import re
import zipfile
from collections import defaultdict
from decimal import Decimal, InvalidOperation
from xml.etree import ElementTree as ET

from openpyxl import load_workbook

NS = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}

BASE = pathlib.Path(r"c:\1c\Cursor_1c\WIM_DEV\bases\WIM_FIn\projects\IMDEV-7330 НовыеТесты")
PROJ = pathlib.Path(
    r"c:\1c\Cursor_1c\WIM_DEV\bases\WIM_FIn\projects"
    r"\IMDEV-9458 Доработка обработки Формирование начислений НДФЛ"
)
OUT_TXT = PROJ / "Тестирование" / "reports" / "avancor10_vs9_report.txt"
OUT_HTML = PROJ / "Документация" / "imdev9458_avancor10_vs9_report.html"

WATCH = (
    "Львов Павел Глебович",
    "Баландин Дмитрий Викторович",
)

RUNS = {
    "9": {
        "uk": BASE / "НДФЛ_Управление_24889_ПоНовому_Аванкор_9.xlsx",
        "pf": BASE / "НДФЛ_Портфели_24889_ПоНовому_Аванкор_9.xlsx",
        "log": BASE / "Лог формирования начисдений НДФЛ_24889_ПоНовому_Аванкор_9.xlsx",
        "msg": BASE / "Сообщения_24889_поНовому_Аванкор_9.txt",
    },
    "10": {
        "uk": BASE / "НДФЛ_Управление_24889_ПоНовому_Аванкор_10.xlsx",
        "pf": BASE / "НДФЛ_Портфели_24889_ПоНовому_Аванкор_10.xlsx",
        "log": BASE / "Лог формирования начисдений НДФЛ_24889_ПоНовому_Аванкор_10.xlsx",
        "msg": BASE / "Сообщения_24889_поНовому_Аванкор_10.txt",
    },
}

MONEY_UK = (
    "СуммаДохода",
    "СуммаВычета",
    "НалогооблагаемаяCумма",
    "СуммаКУдержанию",
    "СуммаРанееУдержанногоНДФЛ",
)
MONEY_PF = (
    "СуммаДохода",
    "СуммаВычета",
    "НалогооблагаемаяCумма",
    "СуммаКУдержанию",
    "СуммаИсчисленногоНалога",
    "СуммаРанееУдержанногоНДФЛ",
)
EPS = Decimal("0.01")


def safe_print(text: str) -> None:
    try:
        print(text)
    except UnicodeEncodeError:
        print(text.encode("ascii", "replace").decode("ascii"))


def to_dec(v) -> Decimal:
    if v is None or v == "":
        return Decimal("0")
    if isinstance(v, Decimal):
        return v
    if isinstance(v, bool):
        return Decimal("0")
    if isinstance(v, int):
        return Decimal(v)
    if isinstance(v, float):
        return Decimal(str(round(v, 4)))
    s = str(v).strip().replace(" ", "").replace("\xa0", "").replace(",", ".")
    if not s:
        return Decimal("0")
    try:
        return Decimal(s)
    except InvalidOperation:
        return Decimal("0")


def money_zero(sums: dict[str, Decimal]) -> bool:
    return (
        sums.get("СуммаДохода", Decimal("0")) == 0
        and sums.get("СуммаКУдержанию", Decimal("0")) == 0
        and sums.get("СуммаИсчисленногоНалога", Decimal("0")) == 0
    )


def cell(row, idx: dict, name: str):
    i = idx.get(name)
    if i is None or i >= len(row):
        return None
    return row[i]


def parse_ndfl(path: pathlib.Path, money_cols: tuple[str, ...]) -> dict:
    safe_print("parse " + path.name)
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    it = ws.iter_rows(values_only=True)
    header = next(it)
    header = [str(c) if c is not None else "" for c in header]
    idx = {name: i for i, name in enumerate(header)}
    docs: dict[str, dict] = {}
    nrows = 0
    for row in it:
        if row is None or all(c is None or str(c).strip() == "" for c in row):
            continue
        nrows += 1
        num = str(cell(row, idx, "НомерДокумента") or "").strip()
        ref = str(cell(row, idx, "Ссылка") or "").strip()
        if not num and ref:
            m = re.search(r"(\d{9,})", ref)
            if m:
                num = m.group(1)
        if not num:
            continue
        cli = str(cell(row, idx, "Клиент") or "").strip()
        port = str(cell(row, idx, "Портфель") or cell(row, idx, "Объект") or "").strip()
        tch = str(cell(row, idx, "ТабличнаяЧасть") or "").strip()
        d = docs.setdefault(
            num,
            {
                "ref": ref,
                "client": cli,
                "ports": set(),
                "tch": set(),
                "rows": 0,
                "sums": {n: Decimal("0") for n in money_cols},
            },
        )
        d["rows"] += 1
        if ref and not d["ref"]:
            d["ref"] = ref
        if cli and not d["client"]:
            d["client"] = cli
        if port:
            d["ports"].add(port)
        if tch:
            d["tch"].add(tch)
        for n in money_cols:
            d["sums"][n] += to_dec(cell(row, idx, n))
    wb.close()
    nonempty = {k: v for k, v in docs.items() if not money_zero(v["sums"])}
    empty = {k: v for k, v in docs.items() if k not in nonempty}
    totals = {n: sum((d["sums"][n] for d in docs.values()), Decimal("0")) for n in money_cols}
    clients = {d["client"] for d in docs.values() if d["client"]}
    return {
        "docs": docs,
        "nonempty": nonempty,
        "empty": empty,
        "totals": totals,
        "clients": clients,
        "rows": nrows,
    }


def parse_log(path: pathlib.Path) -> dict:
    with zipfile.ZipFile(path) as z:
        if "xl/sharedStrings.xml" not in z.namelist():
            strings = []
        else:
            root = ET.fromstring(z.read("xl/sharedStrings.xml"))
            strings = [
                "".join(t.text or "" for t in si.findall(".//m:t", NS))
                for si in root.findall("m:si", NS)
            ]
    header = strings[0] if strings else ""
    err = 0
    m = re.search(r"Ошибок:\s*(\d+)", header)
    if m:
        err = int(m.group(1))
    return {
        "header": header.replace("\n", " | "),
        "errors": err,
        "uk_docs": sum(1 for s in strings if "по управляющей компании" in s),
        "pf_docs": sum(1 for s in strings if "Начисление НДФЛ по портфелю" in s),
    }


def parse_msg(path: pathlib.Path) -> dict:
    text = path.read_text(encoding="utf-8", errors="replace")
    lines = [ln for ln in text.splitlines() if ln.strip()]
    kinds = defaultdict(int)
    for ln in lines:
        if "Невозможно распределить НДФЛ" in ln:
            kinds["distribute"] += 1
        elif "граница актуальности" in ln:
            kinds["border"] += 1
        elif "не проведен" in ln:
            kinds["unposted"] += 1
        elif "отсутствуют данные для расчета" in ln:
            kinds["empty_skip"] += 1
        else:
            kinds["other"] += 1
    return {"lines": len(lines), "kinds": dict(kinds)}


def fmt_n(n) -> str:
    if isinstance(n, Decimal):
        s = f"{n:,.2f}"
    else:
        s = f"{int(n):,}"
    return s.replace(",", " ")


def compare_docs(full: dict, subset: dict, money_cols: tuple[str, ...]) -> dict:
    ka, kb = set(full["docs"]), set(subset["docs"])
    only_full = sorted(ka - kb)
    only_sub = sorted(kb - ka)
    both = sorted(ka & kb)
    diffs = []
    for num in both:
        da, db = full["docs"][num], subset["docs"][num]
        delta = {}
        for n in money_cols:
            dlt = db["sums"][n] - da["sums"][n]
            if abs(dlt) >= EPS:
                delta[n] = dlt
        if delta:
            diffs.append(
                {
                    "num": num,
                    "client": db["client"] or da["client"],
                    "delta": delta,
                }
            )
    return {
        "only_full": only_full,
        "only_sub": only_sub,
        "both": both,
        "diffs": diffs,
        "only_full_nz": [n for n in only_full if n in full["nonempty"]],
        "only_sub_nz": [n for n in only_sub if n in subset["nonempty"]],
    }


def watch_docs(dump: dict) -> dict[str, list[dict]]:
    out = {name: [] for name in WATCH}
    for num, d in dump["docs"].items():
        if d["client"] in out:
            out[d["client"]].append(
                {
                    "num": num,
                    "income": d["sums"].get("СуммаДохода", Decimal("0")),
                    "withhold": d["sums"].get("СуммаКУдержанию", Decimal("0")),
                    "ports": sorted(d["ports"]),
                    "empty": num in dump["empty"],
                }
            )
    return out


def load_run(label: str) -> dict:
    p = RUNS[label]
    safe_print("=== Avancor " + label + " ===")
    return {
        "uk": parse_ndfl(p["uk"], MONEY_UK),
        "pf": parse_ndfl(p["pf"], MONEY_PF),
        "log": parse_log(p["log"]),
        "msg": parse_msg(p["msg"]),
    }


def build_txt(a9, a10, c_uk, c_pf) -> str:
    w9 = watch_docs(a9["uk"])
    w10 = watch_docs(a10["uk"])
    lines = []
    a = lines.append
    a("=== AVANCOR 10 vs 9 (pustoi pul = ne-RDU) ===")
    a("A9 = polnyi nabor. A10 = tot zhe raschet, pul pustoi.")
    a("")
    a("--- Log / messages ---")
    for lab, run in (("9", a9), ("10", a10)):
        a("A{0} log: {1} errors={2}".format(lab, run["log"]["header"][:160], run["log"]["errors"]))
        a("A{0} msg lines={1} kinds={2}".format(lab, run["msg"]["lines"], run["msg"]["kinds"]))
    a("")
    a("--- UK ---")
    for lab, run in (("9", a9), ("10", a10)):
        uk = run["uk"]
        a(
            "A{0} docs={1} empty={2} nonempty={3} clients={4} income={5} withhold={6}".format(
                lab, len(uk["docs"]), len(uk["empty"]), len(uk["nonempty"]),
                len(uk["clients"]), uk["totals"]["СуммаДохода"], uk["totals"]["СуммаКУдержанию"],
            )
        )
    a(
        "UK both={0} only10={1} (nz={2}) only9={3} (nz={4}) money_diffs={5}".format(
            len(c_uk["both"]), len(c_uk["only_sub"]), len(c_uk["only_sub_nz"]),
            len(c_uk["only_full"]), len(c_uk["only_full_nz"]), len(c_uk["diffs"]),
        )
    )
    if c_uk["only_sub_nz"]:
        a("UK nonempty only in A10:")
        for num in c_uk["only_sub_nz"]:
            d = a10["uk"]["docs"][num]
            a("  {0} {1} income={2} withhold={3}".format(
                num, d["client"], d["sums"]["СуммаДохода"], d["sums"]["СуммаКУдержанию"]
            ))
    if c_uk["diffs"]:
        a("UK money diffs (first 15):")
        for item in c_uk["diffs"][:15]:
            a("  {0} {1} {2}".format(item["num"], item["client"], item["delta"]))
    a("")
    a("--- PF ---")
    for lab, run in (("9", a9), ("10", a10)):
        pf = run["pf"]
        a(
            "A{0} docs={1} empty={2} nonempty={3} income={4} withhold={5}".format(
                lab, len(pf["docs"]), len(pf["empty"]), len(pf["nonempty"]),
                pf["totals"]["СуммаДохода"], pf["totals"]["СуммаКУдержанию"],
            )
        )
    a(
        "PF both={0} only10={1} (nz={2}) only9={3} (nz={4}) money_diffs={5}".format(
            len(c_pf["both"]), len(c_pf["only_sub"]), len(c_pf["only_sub_nz"]),
            len(c_pf["only_full"]), len(c_pf["only_full_nz"]), len(c_pf["diffs"]),
        )
    )
    if c_pf["only_sub_nz"]:
        a("PF nonempty only in A10: {0}".format(c_pf["only_sub_nz"][:20]))
    if c_pf["diffs"]:
        a("PF money diffs (first 15):")
        for item in c_pf["diffs"][:15]:
            a("  {0} {1} {2}".format(item["num"], item["client"], item["delta"]))
    a("")
    a("--- Watch ---")
    for name in WATCH:
        a("{0}: A9={1} A10={2}".format(
            name, "YES" if w9[name] else "NO", "YES" if w10[name] else "NO"
        ))
        for d in w10[name]:
            a("  A10 num={0} empty={1} income={2} withhold={3}".format(
                d["num"], d["empty"], d["income"], d["withhold"]
            ))
            if w9[name]:
                d9 = next((x for x in w9[name] if x["num"] == d["num"]), None)
                if d9:
                    a("  A9  num={0} empty={1} income={2} withhold={3}".format(
                        d9["num"], d9["empty"], d9["income"], d9["withhold"]
                    ))
    return "\n".join(lines) + "\n"


def build_html(a9, a10, c_uk, c_pf) -> str:
    def n(x):
        return html.escape(fmt_n(x))

    uk9, uk10 = a9["uk"], a10["uk"]
    pf9, pf10 = a9["pf"], a10["pf"]
    w9 = watch_docs(uk9)
    w10 = watch_docs(uk10)
    ok = (
        len(c_uk["only_sub_nz"]) == 0
        and len(c_uk["diffs"]) == 0
        and len(c_pf["only_sub_nz"]) == 0
        and len(c_pf["diffs"]) == 0
    )
    subset = set(uk10["docs"]) <= set(uk9["docs"]) and set(pf10["docs"]) <= set(pf9["docs"])

    if ok and subset:
        verdict = (
            "По оставшимся клиентам документы и суммы совпали с прогоном 9. "
            "Прогон 10 — подмножество полного набора: пустой пул только отсёк РДУ, расчёт не изменил."
        )
        vcls = "ok"
    elif ok:
        verdict = "Суммы общих документов совпали, но в прогоне 10 есть номера, которых не было в 9."
        vcls = "warn"
    else:
        verdict = "Есть расхождения сумм или лишние живые документы в прогоне 10."
        vcls = "err"

    extra = ""
    for num in c_uk["only_sub_nz"]:
        d = uk10["docs"][num]
        extra += (
            "<tr><td>{0}</td><td>{1}</td><td class='num'>{2}</td><td class='num'>{3}</td></tr>\n"
            .format(html.escape(d["client"]), html.escape(num),
                    n(d["sums"]["СуммаДохода"]), n(d["sums"]["СуммаКУдержанию"]))
        )
    if not extra:
        extra = "<tr><td colspan='4'>Нет. Все живые документы УК прогона 10 есть в прогоне 9.</td></tr>\n"

    diffs = ""
    for item in (c_uk["diffs"] + c_pf["diffs"])[:20]:
        parts = ", ".join("{0}: {1}".format(k, fmt_n(v)) for k, v in item["delta"].items())
        diffs += "<tr><td>{0}</td><td>{1}</td><td>{2}</td></tr>\n".format(
            html.escape(item["num"]), html.escape(item["client"]), html.escape(parts)
        )
    if not diffs:
        diffs = "<tr><td colspan='3'>Расхождений сумм по общим номерам нет.</td></tr>\n"

    watch_rows = ""
    for name in WATCH:
        d10 = w10[name]
        d9 = w9[name]
        if not d10:
            watch_rows += (
                "<tr><td><strong>{}</strong></td><td>{}</td><td colspan='3'>В прогоне 10 нет</td></tr>\n"
                .format(html.escape(name), "был в 9" if d9 else "не было в 9")
            )
            continue
        for d in d10:
            d9m = next((x for x in d9 if x["num"] == d["num"]), None)
            match = ""
            if d9m and d9m["income"] == d["income"] and d9m["withhold"] == d["withhold"]:
                match = "совпало с 9"
            elif d9m:
                match = "номер тот же, суммы другие"
            elif d9:
                match = "в 9 другой номер"
            else:
                match = "в 9 не было"
            watch_rows += (
                "<tr><td><strong>{0}</strong></td><td>{1}</td>"
                "<td class='num'>{2}</td><td class='num'>{3}</td><td>{4}</td></tr>\n"
                .format(html.escape(name), html.escape(d["num"]),
                        n(d["income"]), n(d["withhold"]), html.escape(match))
            )

    return f"""<!DOCTYPE html>
<html lang="ru">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>IMDEV-9458. Сверка прогона 10 с прогоном 9</title>
<style>
  :root {{ --ok:#28a745; --err:#dc3545; --warn:#b78100; --ink:#212529; --soft:#6c757d;
    --bg:#f5f6f8; --card:#fff; --line:#dee2e6; }}
  * {{ box-sizing:border-box; }}
  body {{ margin:0; font-family:"Segoe UI", Arial, Helvetica, sans-serif; font-size:15px;
    line-height:1.45; color:var(--ink); background:var(--bg); }}
  .wrap {{ max-width:960px; margin:0 auto; padding:24px 16px 48px; }}
  header {{ background:#1f4a3a; color:#f7f1e8; border-radius:16px; padding:22px 24px; margin-bottom:20px; }}
  header .kicker {{ font-size:11px; letter-spacing:.1em; text-transform:uppercase; opacity:.75; margin-bottom:6px; }}
  header h1 {{ margin:0 0 8px; font-size:24px; font-family:Georgia, "Times New Roman", serif; }}
  header p {{ margin:0; color:#dce8e2; }}
  section {{ background:var(--card); border:1px solid var(--line); border-radius:12px; padding:18px 20px; margin-bottom:16px; }}
  h2 {{ margin:0 0 12px; font-size:18px; }}
  h3 {{ margin:16px 0 8px; font-size:16px; }}
  p, ul {{ margin:0 0 10px; }}
  ul {{ padding-left:20px; }}
  table {{ width:100%; border-collapse:collapse; margin:10px 0 4px; font-size:14px; }}
  th, td {{ border:1px solid var(--line); padding:8px 10px; text-align:left; vertical-align:top; }}
  th {{ background:#eef2f0; }}
  .num {{ text-align:right; font-variant-numeric:tabular-nums; white-space:nowrap; }}
  .ok {{ color:var(--ok); font-weight:600; }}
  .warn {{ color:var(--warn); font-weight:600; }}
  .err {{ color:var(--err); font-weight:600; }}
  .box-green {{ background:#e8f6ec; border:1px solid #b7e0c2; border-radius:10px; padding:12px 14px; margin:10px 0; }}
  .box-blue {{ background:#e7f4fb; border:1px solid #b9d7ea; border-radius:10px; padding:12px 14px; margin:10px 0; }}
  .grid {{ display:grid; grid-template-columns:repeat(auto-fit,minmax(160px,1fr)); gap:10px; margin:10px 0; }}
  .stat {{ background:#f8faf9; border:1px solid var(--line); border-radius:10px; padding:12px; }}
  .stat .v {{ font-size:22px; font-weight:700; }}
  .stat .l {{ font-size:12px; color:var(--soft); margin-top:4px; }}
  .muted {{ color:var(--soft); font-size:13px; }}
</style>
</head>
<body>
<div class="wrap">
<header>
  <div class="kicker">IMDEV-9458 · тест 10 против теста 9</div>
  <h1>Сверка сумм и документов</h1>
  <p>Прогон 9 — полный набор. Прогон 10 — тот же расчёт, флажок пула включён, поле пула пустое (не-РДУ).</p>
</header>

<section>
  <h2>1. Введение</h2>
  <p>Сверяем не «кто РДУ», а документы и суммы клиентов, которые остались после заполнения с пустым пулом: они должны совпасть с прогоном 9.</p>
  <div class="box-green"><strong>Итог.</strong> <span class="{vcls}">{html.escape(verdict)}</span></div>
</section>

<section>
  <h2>2. Технология</h2>
  <p>Выгрузки логов, сообщений, документов по УК и по портфелю. Строки ТЧ сгруппированы по номеру документа. Пустой документ: доход = 0, к удержанию = 0 (для портфеля ещё исчисленный налог = 0).</p>
  <div class="box-blue">
    <ul>
      <li>Прогон 9: {html.escape(a9["log"]["header"][:200])}</li>
      <li>Прогон 10: {html.escape(a10["log"]["header"][:200])}</li>
    </ul>
  </div>
</section>

<section>
  <h2>3. Статистика</h2>
  <div class="grid">
    <div class="stat"><div class="v">{n(len(uk10["docs"]))}</div><div class="l">УК в прогоне 10</div></div>
    <div class="stat"><div class="v">{n(len(uk9["docs"]))}</div><div class="l">УК в прогоне 9</div></div>
    <div class="stat"><div class="v {'ok' if len(c_uk['diffs'])==0 else 'err'}">{n(len(c_uk['diffs']))}</div><div class="l">Расхождений сумм УК</div></div>
    <div class="stat"><div class="v {'ok' if len(c_pf['diffs'])==0 else 'err'}">{n(len(c_pf['diffs']))}</div><div class="l">Расхождений сумм портфеля</div></div>
  </div>
  <h3>Документы по УК</h3>
  <table>
    <tr><th>Метрика</th><th class="num">Прогон 9</th><th class="num">Прогон 10</th></tr>
    <tr><td>Документов</td><td class="num">{n(len(uk9["docs"]))}</td><td class="num">{n(len(uk10["docs"]))}</td></tr>
    <tr><td>Пустых</td><td class="num">{n(len(uk9["empty"]))}</td><td class="num">{n(len(uk10["empty"]))}</td></tr>
    <tr><td>С суммами</td><td class="num">{n(len(uk9["nonempty"]))}</td><td class="num">{n(len(uk10["nonempty"]))}</td></tr>
    <tr><td>Клиентов</td><td class="num">{n(len(uk9["clients"]))}</td><td class="num">{n(len(uk10["clients"]))}</td></tr>
    <tr><td>Сумма дохода</td><td class="num">{n(uk9["totals"]["СуммаДохода"])}</td><td class="num">{n(uk10["totals"]["СуммаДохода"])}</td></tr>
    <tr><td>Сумма к удержанию</td><td class="num">{n(uk9["totals"]["СуммаКУдержанию"])}</td><td class="num">{n(uk10["totals"]["СуммаКУдержанию"])}</td></tr>
  </table>
  <h3>Документы по портфелю</h3>
  <table>
    <tr><th>Метрика</th><th class="num">Прогон 9</th><th class="num">Прогон 10</th></tr>
    <tr><td>Документов</td><td class="num">{n(len(pf9["docs"]))}</td><td class="num">{n(len(pf10["docs"]))}</td></tr>
    <tr><td>Пустых</td><td class="num">{n(len(pf9["empty"]))}</td><td class="num">{n(len(pf10["empty"]))}</td></tr>
    <tr><td>С суммами</td><td class="num">{n(len(pf9["nonempty"]))}</td><td class="num">{n(len(pf10["nonempty"]))}</td></tr>
    <tr><td>Сумма дохода</td><td class="num">{n(pf9["totals"]["СуммаДохода"])}</td><td class="num">{n(pf10["totals"]["СуммаДохода"])}</td></tr>
    <tr><td>Сумма к удержанию</td><td class="num">{n(pf9["totals"]["СуммаКУдержанию"])}</td><td class="num">{n(pf10["totals"]["СуммаКУдержанию"])}</td></tr>
  </table>
</section>

<section>
  <h2>4. Детали сверки 10 против 9</h2>
  <table>
    <tr><th>Срез</th><th class="num">Общих номеров</th><th class="num">Только в 10, живые</th><th class="num">Только в 9, живые</th><th class="num">Расхождений сумм</th></tr>
    <tr><td>УК</td><td class="num">{n(len(c_uk["both"]))}</td><td class="num">{n(len(c_uk["only_sub_nz"]))}</td><td class="num">{n(len(c_uk["only_full_nz"]))}</td><td class="num">{n(len(c_uk["diffs"]))}</td></tr>
    <tr><td>Портфель</td><td class="num">{n(len(c_pf["both"]))}</td><td class="num">{n(len(c_pf["only_sub_nz"]))}</td><td class="num">{n(len(c_pf["only_full_nz"]))}</td><td class="num">{n(len(c_pf["diffs"]))}</td></tr>
  </table>
  <h3>Живые УК, которых не было в прогоне 9</h3>
  <table>
    <tr><th>Клиент</th><th>Номер</th><th class="num">Сумма дохода</th><th class="num">К удержанию</th></tr>
    {extra}
  </table>
  <h3>Расхождения сумм по общим номерам</h3>
  <table>
    <tr><th>Номер</th><th>Клиент</th><th>Дельта</th></tr>
    {diffs}
  </table>
  <h3>Львов и Баландин</h3>
  <table>
    <tr><th>Клиент</th><th>Номер УК в 10</th><th class="num">Доход</th><th class="num">К удержанию</th><th>С прогоном 9</th></tr>
    {watch_rows}
  </table>
</section>

<section>
  <h2>5. Возможности</h2>
  <ul>
    <li>Зеркальный прогон с пулом «Розничное ДУ».</li>
    <li>Если появится дельта по общему номеру — открыть документ в базе 24889 и сравнить ТЧ.</li>
  </ul>
</section>

<section>
  <h2>6. Выводы</h2>
  <ol>
    <li>Прогон 10 короче полного набора 9: остались клиенты с пустым пулом.</li>
    <li>По общим документам суммы {"совпали" if ok else "разошлись"}.</li>
    <li>Ошибок в логе прогона 10: {n(a10["log"]["errors"])}.</li>
  </ol>
</section>
<p class="muted">Отчёт по выгрузкам прогонов 9 и 10, задача IMDEV-9458.</p>
</div>
</body>
</html>
"""


def main() -> int:
    a9 = load_run("9")
    a10 = load_run("10")
    c_uk = compare_docs(a9["uk"], a10["uk"], MONEY_UK)
    c_pf = compare_docs(a9["pf"], a10["pf"], MONEY_PF)
    txt = build_txt(a9, a10, c_uk, c_pf)
    OUT_TXT.parent.mkdir(parents=True, exist_ok=True)
    OUT_TXT.write_text(txt, encoding="utf-8")
    OUT_HTML.write_text(build_html(a9, a10, c_uk, c_pf), encoding="utf-8")
    safe_print(txt)
    safe_print("TXT OK")
    safe_print("HTML OK")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
