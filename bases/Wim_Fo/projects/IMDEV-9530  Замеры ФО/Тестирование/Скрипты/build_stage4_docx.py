#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Сборка DOCX-отчета этапа 4 из HTML-отчета тестирования замеров IMDEV-9530.
"""
import base64
import io
import re
from pathlib import Path

from bs4 import BeautifulSoup, NavigableString
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
HTML_PATH = ROOT / "Отчет о тестировании замеров, этап 4.html"
DOCX_PATH = ROOT / "Отчет о тестировании замеров, этап 4.docx"

INK = RGBColor(0x18, 0x21, 0x2D)
MUTED = RGBColor(0x5D, 0x6B, 0x7C)
HEAD = "1D3557"
OK = RGBColor(0x1F, 0x9D, 0x55)
BAD = RGBColor(0xD6, 0x45, 0x45)
ACCENT = RGBColor(0x2A, 0x6F, 0xDB)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

PAGE_WIDTH_CM = 29.7
PAGE_HEIGHT_CM = 21.0
MARGIN_CM = 1.25
CONTENT_CM = PAGE_WIDTH_CM - 2 * MARGIN_CM


def cm_to_dxa(value_cm):
    return int(round(value_cm / 2.54 * 1440))


def set_run_font(run, name, size_pt, bold=False, color=None, italic=False):
    run.bold = bold
    run.italic = italic
    run.font.name = name
    run.font.size = Pt(size_pt)
    if color is not None:
        run.font.color.rgb = color
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), name)
    rfonts.set(qn("w:hAnsi"), name)
    rfonts.set(qn("w:cs"), name)


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    for old in tc_pr.findall(qn("w:shd")):
        tc_pr.remove(old)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_width(cell, dxa):
    tc_pr = cell._tc.get_or_add_tcPr()
    for old in tc_pr.findall(qn("w:tcW")):
        tc_pr.remove(old)
    tc_w = OxmlElement("w:tcW")
    tc_w.set(qn("w:w"), str(dxa))
    tc_w.set(qn("w:type"), "dxa")
    tc_pr.append(tc_w)


def set_cell_borders(cell, color="DBE2EA", sz="4", left_color=None, left_sz=None):
    tc_pr = cell._tc.get_or_add_tcPr()
    for old in tc_pr.findall(qn("w:tcBorders")):
        tc_pr.remove(old)
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement("w:" + edge)
        el.set(qn("w:val"), "single")
        if edge == "left" and left_color:
            el.set(qn("w:sz"), left_sz or "16")
            el.set(qn("w:color"), left_color)
        else:
            el.set(qn("w:sz"), sz)
            el.set(qn("w:color"), color)
        el.set(qn("w:space"), "0")
        borders.append(el)
    tc_pr.append(borders)


def set_paragraph_border(paragraph, color, fill=None):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "18")
    left.set(qn("w:space"), "8")
    left.set(qn("w:color"), color)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "nil")
    p_bdr.append(left)
    p_pr.append(p_bdr)
    if fill:
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), fill)
        p_pr.append(shd)


def set_paragraph_spacing(paragraph, before=0, after=6, line=240):
    pf = paragraph.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line / 240.0


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant = OxmlElement("w:cantSplit")
    tr_pr.append(cant)


def mark_header_row(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    tr_pr.append(header)


def set_table_fixed(table, widths_dxa):
    table.autofit = False
    table.allow_autofit = False
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    total = sum(widths_dxa)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")
    mar = tbl_pr.find(qn("w:tblCellMar"))
    if mar is not None:
        tbl_pr.remove(mar)
    mar = OxmlElement("w:tblCellMar")
    for edge, value in (("top", "40"), ("left", "70"), ("bottom", "40"), ("right", "70")):
        el = OxmlElement("w:" + edge)
        el.set(qn("w:w"), value)
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tbl_pr.append(mar)
    grid = tbl.find(qn("w:tblGrid"))
    if grid is not None:
        tbl.remove(grid)
    grid = OxmlElement("w:tblGrid")
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    tbl.insert(1, grid)


def classes_of(tag):
    raw = tag.get("class") or []
    if isinstance(raw, str):
        return raw.split()
    return list(raw)


def inline_pieces(node):
    pieces = []
    for child in node.children:
        if isinstance(child, NavigableString):
            text = re.sub(r"\s+", " ", str(child))
            if text:
                pieces.append(("text", text, None))
        elif getattr(child, "name", None) == "br":
            pieces.append(("br", "", None))
        elif child.name in ("b", "strong"):
            text = re.sub(r"\s+", " ", child.get_text())
            pieces.append(("bold", text, classes_of(child)))
        elif child.name == "code":
            text = re.sub(r"\s+", " ", child.get_text()).strip()
            if text:
                pieces.append(("code", text, None))
        else:
            pieces.extend(inline_pieces(child))
    while pieces and pieces[0][0] == "text":
        trimmed = pieces[0][1].lstrip()
        if trimmed:
            pieces[0] = ("text", trimmed, None)
            break
        pieces.pop(0)
    while pieces and pieces[-1][0] == "text":
        trimmed = pieces[-1][1].rstrip()
        if trimmed:
            pieces[-1] = ("text", trimmed, None)
            break
        pieces.pop()
    return pieces


def add_pieces(paragraph, pieces, size_pt, base_color=INK):
    for kind, text, extra in pieces:
        if kind == "br":
            paragraph.add_run().add_break()
            continue
        if not text:
            continue
        run = paragraph.add_run(text)
        if kind == "code":
            set_run_font(run, "Consolas", max(size_pt - 0.5, 8), color=RGBColor(0x1D, 0x35, 0x57))
        elif kind == "bold":
            color = base_color
            if extra and "yes" in extra:
                color = OK
            elif extra and "no" in extra:
                color = BAD
            set_run_font(run, "Calibri", size_pt, bold=True, color=color)
        else:
            set_run_font(run, "Calibri", size_pt, color=base_color)


def add_inline_paragraph(document, node, size_pt=11, before=0, after=8, color=INK, italic=False):
    paragraph = document.add_paragraph()
    set_paragraph_spacing(paragraph, before, after, 252)
    add_pieces(paragraph, inline_pieces(node), size_pt, color)
    if italic:
        for run in paragraph.runs:
            run.italic = True
    return paragraph


def widths_for(headers):
    key = tuple(headers)
    presets = {
        (
            "Внешняя обработка",
            "Файл",
            "Версия: оригинал -> с замерами",
            "Строка списка ключевых операций",
            "Ключи замеров (общая запись и шаги)",
            "Точки входа с замером",
        ): [3.4, 4.2, 2.4, 4.2, 6.6, 6.5],
        (
            "Сценарий",
            "Обработка",
            "Точка входа",
            "Ключ",
            "Ожидалось",
            "Факт",
            "Совпало",
        ): [3.3, 3.0, 3.5, 8.5, 5.2, 2.2, 1.6],
        (
            "Что сравнивалось",
            "Объем",
            "Итог",
            "Подробно",
        ): [8.4, 6.2, 2.6, 10.1],
        (
            "Прогон",
            "Факт",
            "Итог",
            "Подробно",
        ): [6.6, 4.6, 2.4, 13.7],
        (
            "Версия",
            "Шаг",
            "Сценарий",
            "Длительность, с",
            "Ошибка точки входа",
            "Записей замеров",
            "Примечание",
        ): [1.7, 3.5, 6.4, 2.2, 6.4, 1.8, 5.3],
        (
            "Начало замера",
            "Сеанс",
            "Ключ",
            "Время, с",
            "Вес",
            "Ошибка",
            "Комментарий",
        ): [3.2, 1.3, 8.5, 1.6, 1.2, 1.4, 10.1],
    }
    if key in presets:
        values = presets[key]
    else:
        values = [CONTENT_CM / max(len(headers), 1)] * len(headers)
    scale = CONTENT_CM / sum(values)
    return [cm_to_dxa(v * scale) for v in values]


def fill_down_scenario(rows):
    """Повторяет сценарий в пустых строках сверки, чтобы строка читалась отдельно."""
    if not rows:
        return rows
    headers = [piece_text(cell) for cell in rows[0]]
    if headers[:3] != ["Сценарий", "Обработка", "Точка входа"]:
        return rows
    last = ["", "", ""]
    filled = [rows[0]]
    for row in rows[1:]:
        new_row = []
        for index, cell in enumerate(row):
            text = piece_text(cell).strip()
            inherited = False
            if index < 1 and not text:
                text = last[index]
                inherited = True
                pieces = [("text", text, None)] if text else []
            else:
                pieces = cell["pieces"]
            if index < 1 and text:
                last[index] = text
            new_row.append({
                "pieces": pieces,
                "numeric": cell["numeric"],
                "inherited": inherited,
            })
        filled.append(new_row)
    return filled


def piece_text(cell):
    return "".join(text for _kind, text, _extra in cell["pieces"]).strip()


def parse_row(tr):
    cells = []
    for td in tr.find_all(["th", "td"], recursive=False):
        cells.append({
            "pieces": inline_pieces(td),
            "numeric": "n" in classes_of(td),
            "inherited": False,
        })
    return cells


def add_table(document, table_tag):
    raw_rows = []
    row_meta = []
    for tr in table_tag.find_all("tr", recursive=False):
        raw_rows.append(parse_row(tr))
        row_meta.append(classes_of(tr))
    if not raw_rows:
        return
    headers = [piece_text(cell) for cell in raw_rows[0]]
    body = fill_down_scenario(raw_rows)
    # fill_down returns header + data; row_meta[0] is header
    if body is not raw_rows:
        data_meta = row_meta[1:]
    else:
        data_meta = row_meta[1:]
    widths = widths_for(headers)
    table = document.add_table(rows=len(body), cols=len(headers))
    table.style = "Table Grid"
    set_table_fixed(table, widths)
    dense = len(body) > 8
    font_size = 8 if dense else 10
    for r_index, row_cells in enumerate(body):
        row = table.rows[r_index]
        prevent_row_split(row)
        is_header = r_index == 0
        if is_header:
            mark_header_row(row)
        meta = [] if is_header else data_meta[r_index - 1]
        left_color = None
        if "ok" in meta:
            left_color = "1F9D55"
        elif "bad" in meta:
            left_color = "D64545"
        zebra = (not is_header) and (r_index % 2 == 0)
        for c_index, cell_info in enumerate(row_cells):
            cell = row.cells[c_index]
            set_cell_width(cell, widths[c_index])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            if is_header:
                shade_cell(cell, HEAD)
                set_cell_borders(cell, color="1D3557", sz="4")
            else:
                if zebra:
                    shade_cell(cell, "F4F7FB")
                set_cell_borders(
                    cell,
                    left_color=left_color if c_index == 0 else None,
                    left_sz="16",
                )
            paragraph = cell.paragraphs[0]
            set_paragraph_spacing(paragraph, 0, 0, 230)
            if cell_info["numeric"] and not is_header:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            color = WHITE if is_header else (MUTED if cell_info.get("inherited") else INK)
            pieces = cell_info["pieces"]
            if is_header:
                pieces = [("bold", piece_text(cell_info), None)]
            add_pieces(paragraph, pieces, font_size, color)
            if is_header:
                for run in paragraph.runs:
                    set_run_font(run, "Calibri", font_size, bold=True, color=WHITE)
    document.add_paragraph().paragraph_format.space_after = Pt(4)


def add_cards(document, cards_div):
    cards = cards_div.find_all("div", class_="card")
    if not cards:
        return
    cols = 3
    rows = (len(cards) + cols - 1) // cols
    table = document.add_table(rows=rows, cols=cols)
    table.style = "Table Grid"
    width = cm_to_dxa(CONTENT_CM / cols)
    set_table_fixed(table, [width] * cols)
    for index, card in enumerate(cards):
        row = table.rows[index // cols]
        prevent_row_split(row)
        cell = row.cells[index % cols]
        set_cell_width(cell, width)
        shade_cell(cell, "F7F9FC")
        set_cell_borders(cell, color="DBE2EA", sz="4")
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        number = card.find("b")
        caption = card.find("span")
        p1 = cell.paragraphs[0]
        set_paragraph_spacing(p1, 2, 0, 230)
        p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p1.add_run(number.get_text(strip=True) if number else "")
        set_run_font(run, "Calibri", 20, bold=True, color=ACCENT)
        p2 = cell.add_paragraph()
        set_paragraph_spacing(p2, 0, 2, 220)
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run2 = p2.add_run(caption.get_text(" ", strip=True) if caption else "")
        set_run_font(run2, "Calibri", 9, color=MUTED)
    # empty trailing cells
    for index in range(len(cards), rows * cols):
        cell = table.rows[index // cols].cells[index % cols]
        set_cell_width(cell, width)
        shade_cell(cell, "FFFFFF")
        set_cell_borders(cell, color="FFFFFF", sz="4")
    spacer = document.add_paragraph()
    set_paragraph_spacing(spacer, 2, 2, 200)


def add_note(document, note):
    paragraph = document.add_paragraph()
    set_paragraph_spacing(paragraph, 2, 8, 246)
    set_paragraph_border(paragraph, "C98A04", "FFF8EC")
    add_pieces(paragraph, inline_pieces(note), 11, INK)


def add_figure(document, figure):
    img = figure.find("img")
    caption = figure.find("figcaption")
    if img and img.get("src", "").startswith("data:image"):
        payload = img["src"].split(",", 1)[1]
        data = base64.b64decode(payload)
        stream = io.BytesIO(data)
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_paragraph_spacing(paragraph, 6, 2, 230)
        run = paragraph.add_run()
        run.add_picture(stream, width=Cm(24.5))
    if caption:
        cap = document.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_paragraph_spacing(cap, 0, 10, 230)
        run = cap.add_run(caption.get_text(" ", strip=True))
        set_run_font(run, "Calibri", 10, italic=True, color=MUTED)


def add_list(document, ul):
    for li in ul.find_all("li", recursive=False):
        paragraph = document.add_paragraph(style="List Bullet")
        set_paragraph_spacing(paragraph, 1, 3, 240)
        add_pieces(paragraph, inline_pieces(li), 11, INK)


def configure_styles(document):
    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = INK
    rpr = normal.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), "Calibri")
    rfonts.set(qn("w:hAnsi"), "Calibri")
    rfonts.set(qn("w:cs"), "Calibri")
    lang = rpr.find(qn("w:lang"))
    if lang is None:
        lang = OxmlElement("w:lang")
        rpr.append(lang)
    lang.set(qn("w:val"), "ru-RU")
    for style_name, size, before, after, color in (
        ("Heading 1", 18, 0, 8, RGBColor(0x1D, 0x35, 0x57)),
        ("Heading 2", 14, 14, 6, RGBColor(0x1D, 0x35, 0x57)),
    ):
        style = document.styles[style_name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        style.font.italic = False
        pf = style.paragraph_format
        pf.space_before = Pt(before)
        pf.space_after = Pt(after)
        pf.line_spacing = 1.08


def add_page_field(paragraph, instruction):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(separate)
    run._r.append(end)
    set_run_font(run, "Calibri", 9, color=MUTED)


def configure_section(document):
    section = document.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Cm(PAGE_WIDTH_CM)
    section.page_height = Cm(PAGE_HEIGHT_CM)
    section.left_margin = Cm(MARGIN_CM)
    section.right_margin = Cm(MARGIN_CM)
    section.top_margin = Cm(1.4)
    section.bottom_margin = Cm(1.3)
    section.header_distance = Cm(0.4)
    section.footer_distance = Cm(0.4)
    header = section.header
    header.is_linked_to_previous = False
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_paragraph_spacing(hp, 0, 2, 200)
    run = hp.add_run("IMDEV-9530  |  этап 4  |  тестирование замеров")
    set_run_font(run, "Calibri", 9, color=MUTED)
    p_pr = hp._p.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "DBE2EA")
    p_bdr.append(bottom)
    p_pr.append(p_bdr)
    footer = section.footer
    footer.is_linked_to_previous = False
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_paragraph_spacing(fp, 2, 0, 200)
    run = fp.add_run("Стр. ")
    set_run_font(run, "Calibri", 9, color=MUTED)
    add_page_field(fp, " PAGE ")
    run = fp.add_run(" из ")
    set_run_font(run, "Calibri", 9, color=MUTED)
    add_page_field(fp, " NUMPAGES ")


def build():
    html = HTML_PATH.read_text(encoding="utf-8")
    soup = BeautifulSoup(html, "html.parser")
    wrap = soup.select_one("div.wrap")
    document = Document()
    configure_styles(document)
    configure_section(document)
    document.core_properties.title = "Тестирование замеров внешних обработок лимитов и поручений (IMDEV-9530, этап 4)"
    document.core_properties.subject = "IMDEV-9530, этап 4"
    document.core_properties.category = "Отчет о тестировании"

    for child in wrap.children:
        if isinstance(child, NavigableString):
            continue
        name = child.name
        cls = classes_of(child)
        if name == "h1":
            paragraph = document.add_paragraph(child.get_text(" ", strip=True), style="Heading 1")
            set_paragraph_spacing(paragraph, 0, 8, 240)
        elif name == "h2":
            paragraph = document.add_paragraph(child.get_text(" ", strip=True), style="Heading 2")
            set_paragraph_spacing(paragraph, 12, 4, 240)
        elif name == "p":
            add_inline_paragraph(document, child, 11, 2, 8, INK)
        elif name == "div" and "verdict" in cls:
            paragraph = document.add_paragraph()
            set_paragraph_spacing(paragraph, 4, 8, 246)
            set_paragraph_border(paragraph, "1F9D55", "E9F7EF")
            pieces = inline_pieces(child)
            if pieces and pieces[0][0] == "text":
                pieces[0] = ("bold", pieces[0][1], None)
            else:
                pieces.insert(0, ("bold", child.get_text(" ", strip=True), None))
                pieces = pieces[:1]
            add_pieces(paragraph, pieces, 11, OK)
        elif name == "div" and "cards" in cls:
            add_cards(document, child)
        elif name == "div" and "note" in cls:
            add_note(document, child)
        elif name == "div" and "shots" in cls:
            for figure in child.find_all("figure", recursive=False):
                add_figure(document, figure)
        elif name == "div":
            table = child.find("table", recursive=False)
            if table:
                add_table(document, table)
        elif name == "ul":
            add_list(document, child)
        elif name == "figure":
            add_figure(document, child)

    document.save(DOCX_PATH)
    print("saved", DOCX_PATH.name)
    print("bytes", DOCX_PATH.stat().st_size)


if __name__ == "__main__":
    build()
