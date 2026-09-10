#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Analyze OK Demand test data structure."""

from pathlib import Path
from openpyxl import load_workbook

p = Path(r"c:\Users\Acer\Downloads\OK_test_Demand") / "Тест_Деманд.xlsx"
wb = load_workbook(p, data_only=True, read_only=True)
ws = wb["Массив"]

week_starts = [13, 22, 31, 40, 49]
rows = []
for row in ws.iter_rows(min_row=8, max_row=435, max_col=76, values_only=True):
    code = row[0]
    if not code:
        continue
    stock = float(row[5] or 0)
    prod_m = float(row[6] or 0)
    prod_w = [float(row[c - 1] or 0) for c in range(8, 13)]
    name = row[1]
    su = row[2]
    cat = row[3]
    atype = row[4]
    fc_weeks = []
    for ws_ in week_starts:
        ch = [float(row[ws_ + i - 1] or 0) for i in range(8)]
        tot = float(row[ws_ + 8 - 1] or 0)
        fc_weeks.append((ch, tot))
    fc_m = [float(row[58 + i - 1] or 0) for i in range(8)]
    fc_m_tot = float(row[65] or 0)
    rows.append(
        dict(
            code=code,
            name=name,
            su=su,
            cat=cat,
            atype=atype,
            stock=stock,
            prod_m=prod_m,
            prod_w=prod_w,
            fc_weeks=fc_weeks,
            fc_m=fc_m,
            fc_m_tot=fc_m_tot,
        )
    )

print("n rows", len(rows))
supply = sum(r["stock"] + sum(r["prod_w"]) for r in rows)
demand = sum(sum(ch) for r in rows for ch, _ in r["fc_weeks"])
print("total supply open+prod", round(supply, 1))
print("total demand 5w channels", round(demand, 1))
print("sum fc month col66", round(sum(r["fc_m_tot"] for r in rows), 1))
print("sum fc month channels", round(sum(sum(r["fc_m"]) for r in rows), 1))
print("target total", 20163.5)

deficits = 0
surplus = 0
for r in rows:
    avail = r["stock"] + sum(r["prod_w"])
    dem = sum(sum(ch) for ch, _ in r["fc_weeks"])
    if avail + 1e-6 < dem:
        deficits += 1
    elif avail > dem + 1e-6:
        surplus += 1
print("deficit SKU", deficits, "surplus SKU", surplus)

sus = set(str(r["su"]) for r in rows)
print("SU samples", list(sus)[:20], "n unique", len(sus))
print("named", sum(1 for r in rows if r["name"]))

bad = 0
for r in rows[:50]:
    for ch, tot in r["fc_weeks"]:
        if abs(sum(ch) - tot) > 1:
            bad += 1
print("week total mismatch samples", bad)

ch_names = ["FS_promo", "FS_reg", "RDC", "FT", "EK", "KPi", "Ecom", "RKP"]
ch_sum = [0.0] * 8
for r in rows:
    for i in range(8):
        ch_sum[i] += r["fc_m"][i]
print("fc month by channel:")
for n, s in zip(ch_names, ch_sum):
    print(n, round(s, 1))
print("sum", round(sum(ch_sum), 1))

# weekly simulation rough deficit
weekly_def = 0
for r in rows:
    carry = r["stock"]
    for w in range(5):
        avail = carry + r["prod_w"][w]
        dem = sum(r["fc_weeks"][w][0])
        if avail + 1e-9 < dem:
            weekly_def += 1
            carry = 0
        else:
            carry = avail - dem
print("week-sku deficit events", weekly_def)

# units hint: maybe data in kg, target in tons
print("demand/1000", round(demand / 1000, 1))
print("fc_m_tot/1000", round(sum(r["fc_m_tot"] for r in rows) / 1000, 1))

wb.close()
