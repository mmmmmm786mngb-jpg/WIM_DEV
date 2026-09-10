#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Copy original Test_Demand.xlsx and fill red 'Plan mes' columns on sheet Massiv
with Excel FORMULAS. Intermediate weekly logic lives on sheet 'Расчёт' (also formulas).
"""

from pathlib import Path
from shutil import copy2

from openpyxl import load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

SRC = Path(r"c:\Users\Acer\Downloads\OK_test_Demand") / "Тест_Деманд.xlsx"
OUT = Path(r"c:\Users\Acer\Downloads\OK_test_Demand") / "Тест_Деманд_ФОРМУЛЫ_Массив.xlsx"

# Channels in weekly forecast blocks (col offset 0..7)
CH_RU = ["ФС промо", "ФС рег", "РДЦ", "ФТ", "ЭК", "КПи", "E-commerce", "РКП"]
# Priority when deficit: promo, reg, RDC, RKP, EK, FT, KPi, Ecom
PRIORITY = [0, 1, 2, 7, 4, 3, 5, 6]

# Massiv columns
STOCK_COL = 6          # F
PROD_COLS = [8, 9, 10, 11, 12]  # H-L weeks 1-5
WEEK_FC_STARTS = [13, 22, 31, 40, 49]  # M, V, AE, AN, AW
# Plan mes red columns
PLAN_FS = 67           # BO - FS total (promo+reg)
PLAN_CH_START = 68     # BP .. BW = 8 channels
PLAN_TOTAL = 76        # BX
TARGET_ROW = 5         # tons in BP5:BW5
DATA_START = 8

FILL_NOTE = PatternFill("solid", fgColor="FFF2CC")
FILL_CALC = PatternFill("solid", fgColor="C6EFCE")
FONT_B = Font(bold=True, name="Calibri", size=11)
FONT_N = Font(name="Calibri", size=9)
THIN = Border(
    left=Side(style="thin", color="B0B0B0"),
    right=Side(style="thin", color="B0B0B0"),
    top=Side(style="thin", color="B0B0B0"),
    bottom=Side(style="thin", color="B0B0B0"),
)


def find_last_data_row(ws):
    last = DATA_START
    for r in range(DATA_START, 2000):
        if ws.cell(r, 1).value not in (None, ""):
            last = r
        elif r > last + 5:
            break
    return last


def main():
    print("Copy source...")
    if OUT.exists():
        OUT.unlink()
    copy2(SRC, OUT)

    print("Open copy (may take ~1 min)...")
    wb = load_workbook(OUT)
    ws = wb["Массив"]
    last = find_last_data_row(ws)
    print("Data rows", DATA_START, "..", last, "n=", last - DATA_START + 1)

    # ----- helper sheet Расчёт -----
    if "Расчёт" in wb.sheetnames:
        del wb["Расчёт"]
    calc = wb.create_sheet("Расчёт", 0)

    # Header row 7 to align with Massiv row numbers (rows 1-7 titles)
    calc["A1"] = "ПРОМЕЖУТОЧНЫЙ РАСЧЁТ (все ячейки с данными — ФОРМУЛЫ, ссылаются на лист Массив)"
    calc["A1"].font = Font(bold=True, size=12, color="1F4E79")
    calc["A1"].fill = FILL_NOTE
    calc.merge_cells("A1:P1")

    calc["A2"] = (
        "Приоритет при дефиците: ФС промо -> ФС рег -> РДЦ -> РКП -> ЭК -> ФТ -> КПи -> E-commerce. "
        "Недельный перенос остатка. План мес на Массиве (красные BO:BX) = ссылки на Итог_* ниже."
    )
    calc["A2"].fill = FILL_NOTE
    calc.merge_cells("A2:P2")

    calc["A3"] = "Цели каналов, кг (из Массив строка 5, тонны*1000):"
    for i, name in enumerate(CH_RU):
        col_m = PLAN_CH_START + i
        calc.cell(4, 1 + i, name)
        # formula: target tons * 1000
        calc.cell(5, 1 + i, f"=Массив!{get_column_letter(col_m)}$5*1000")
        calc.cell(5, 1 + i).fill = FILL_CALC
        calc.cell(5, 1 + i).font = FONT_N

    calc["A6"] = "Строки 8+ синхронизированы с Массив. Не меняйте номера строк."
    calc["A6"].fill = FILL_NOTE

    # Column layout on Расчёт starting row 7 headers, data from row 8:
    # A Kod
    # B Dostupno
    # C Spros
    # D Deficit
    # E Status
    # F-M Syroy ch0-7
    # N Syroy itogo
    # O-V Itog ch0-7  (FINAL -> linked from Massiv)
    # W Itog itogo
    # Then weekly detail from col 24 (X)

    h = [
        "Код",
        "Доступно",
        "Спрос",
        "Дефицит",
        "Статус",
    ]
    for name in CH_RU:
        h.append(f"Сырой_{name}")
    h.append("Сырой_Итого")
    for name in CH_RU:
        h.append(f"Итог_{name}")
    h.append("Итог_Итого")

    week_start_col = 24  # X
    for w in range(1, 6):
        h.append(f"Н{w}_Доступно")
        for step, ch_i in enumerate(PRIORITY):
            h.append(f"Н{w}_План_{CH_RU[ch_i]}")
            h.append(f"Н{w}_Ост{step+1}")

    for c, name in enumerate(h, 1):
        cell = calc.cell(7, c, name)
        cell.font = Font(bold=True, color="FFFFFF", size=8)
        cell.fill = PatternFill("solid", fgColor="305496")
        cell.alignment = Alignment(wrap_text=True, horizontal="center")
    calc.row_dimensions[7].height = 30

    n_rows = last
    week_plan_cols_template = None  # built once from col numbers

    for r in range(DATA_START, last + 1):
        # A code
        calc.cell(r, 1, f"=Массив!A{r}")
        # B supply = F + H+I+J+K+L
        calc.cell(
            r,
            2,
            f"=Массив!F{r}+Массив!H{r}+Массив!I{r}+Массив!J{r}+Массив!K{r}+Массив!L{r}",
        )
        # C demand = sum of all weekly channel forecasts (not weekly totals to avoid double)
        # M:T without U, V:AC without AD, etc. — use SUM of each week 8 channels
        parts = []
        for ws_ in WEEK_FC_STARTS:
            a = get_column_letter(ws_)
            b = get_column_letter(ws_ + 7)
            parts.append(f"SUM(Массив!{a}{r}:{b}{r})")
        calc.cell(r, 3, "=" + "+".join(parts))
        calc.cell(r, 4, f"=MAX(0,C{r}-B{r})")
        # Status uses final leftover approx
        calc.cell(
            r,
            5,
            f'=IF(D{r}>0.0001,"ДЕФИЦИТ",IF(B{r}-W{r}>0.0001,"ПРОФИЦИТ","БАЛАНС"))',
        )

        # Weekly blocks
        col = week_start_col
        week_plan_cols = {ch: [] for ch in range(8)}

        for w in range(5):
            if w == 0:
                avail_f = f"=Массив!F{r}+Массив!H{r}"
            else:
                prev_end = week_start_col + (w - 1) * 17 + 16
                prod_letter = get_column_letter(PROD_COLS[w])
                avail_f = f"={get_column_letter(prev_end)}{r}+Массив!{prod_letter}{r}"
            calc.cell(r, col, avail_f)
            remain_col = col
            col += 1
            for step, ch_i in enumerate(PRIORITY):
                fc_letter = get_column_letter(WEEK_FC_STARTS[w] + ch_i)
                plan_f = f"=MIN(Массив!{fc_letter}{r},{get_column_letter(remain_col)}{r})"
                calc.cell(r, col, plan_f)
                week_plan_cols[ch_i].append(col)
                plan_col = col
                col += 1
                rem_f = f"={get_column_letter(remain_col)}{r}-{get_column_letter(plan_col)}{r}"
                calc.cell(r, col, rem_f)
                remain_col = col
                col += 1

        if week_plan_cols_template is None:
            week_plan_cols_template = {k: list(v) for k, v in week_plan_cols.items()}

        # Raw month F=6 .. M=13 -> cols 6..13 on calc (1-based: 6=F ... 13=M)
        for ch_i in range(8):
            refs = "+".join(f"{get_column_letter(c)}{r}" for c in week_plan_cols[ch_i])
            calc.cell(r, 6 + ch_i, f"={refs}")
        calc.cell(r, 14, f"=SUM(F{r}:M{r})")

        # Final O=15 .. V=22 with cap vs targets on row 5 of calc (A5:H5)
        for ch_i in range(8):
            raw_l = get_column_letter(6 + ch_i)
            tgt_l = get_column_letter(1 + ch_i)  # A5..H5
            calc.cell(
                r,
                15 + ch_i,
                f"=IFERROR({raw_l}{r}*MIN(1,{tgt_l}$5/SUM({raw_l}${DATA_START}:{raw_l}${last})),0)",
            )
        calc.cell(r, 23, f"=SUM(O{r}:V{r})")

        # light fill for formula cells
        for c in range(1, 24):
            calc.cell(r, c).fill = FILL_CALC
            calc.cell(r, c).font = FONT_N
        for c in range(week_start_col, week_start_col + 5 * 17):
            calc.cell(r, c).fill = FILL_CALC
            calc.cell(r, c).font = FONT_N

    calc.freeze_panes = "B8"
    calc.auto_filter.ref = f"A7:W{last}"
    calc.column_dimensions["A"].width = 12
    calc.column_dimensions["E"].width = 11

    # ----- fill Massiv red Plan mes with formulas -----
    print("Write Plan mes formulas on Massiv...")
    comment = Comment(
        "Формула: итог после недельного приоритета и урезки до лимита канала. "
        "Промежуточный расчёт — лист «Расчёт».",
        "Reshenie",
    )
    comment.width = 280
    comment.height = 60

    for r in range(DATA_START, last + 1):
        # BO FS = promo + reg
        ws.cell(r, PLAN_FS).value = f"=BP{r}+BQ{r}"
        # BP..BW = Итог channels O..V on Расчёт (cols 15..22)
        for i in range(8):
            calc_col = get_column_letter(15 + i)
            cell = ws.cell(r, PLAN_CH_START + i)
            cell.value = f"=Расчёт!{calc_col}{r}"
        # BX total
        ws.cell(r, PLAN_TOTAL).value = f"=SUM(BP{r}:BW{r})"

    # header comments
    ws.cell(7, PLAN_CH_START).comment = comment

    # row 4 check sums (formulas)
    ws.cell(4, PLAN_FS - 1).value = "Проверка СУММ плана (кг):"
    ws.cell(4, PLAN_FS - 1).font = FONT_B
    ws.cell(4, PLAN_FS - 1).fill = FILL_NOTE
    for i in range(8):
        col = PLAN_CH_START + i
        letter = get_column_letter(col)
        ws.cell(4, col).value = f"=SUM({letter}{DATA_START}:{letter}{last})"
        ws.cell(4, col).fill = FILL_CALC
        ws.cell(4, col).font = FONT_N
    ws.cell(4, PLAN_TOTAL).value = f"=SUM({get_column_letter(PLAN_TOTAL)}{DATA_START}:{get_column_letter(PLAN_TOTAL)}{last})"
    ws.cell(4, PLAN_TOTAL).fill = FILL_CALC

    # Optional: month forecast BF:BN as sum of weeks (formulas) — helps transparency
    # BF=58 FS promo month = M+V+AE+AN+AW etc.
    print("Fill month forecast helper formulas BF:BN...")
    for r in range(DATA_START, last + 1):
        for ch_i in range(8):
            refs = "+".join(
                f"{get_column_letter(WEEK_FC_STARTS[w] + ch_i)}{r}" for w in range(5)
            )
            ws.cell(r, 58 + ch_i).value = f"={refs}"
        ws.cell(r, 66).value = f"=SUM(BF{r}:BM{r})"

    # Readme sheet
    if "00_Как_сдавать" in wb.sheetnames:
        del wb["00_Как_сдавать"]
    readme = wb.create_sheet("00_Как_сдавать", 0)
    lines = [
        "ЧТО СДАВАТЬ",
        "Заполнены КРАСНЫЕ колонки «План мес» на листе Массив (BO:BX) — ФОРМУЛАМИ.",
        "Промежуточные расчёты — лист «Расчёт» (тоже формулы, строки = строки Массива).",
        "",
        "КАК ПРОВЕРИТЬ",
        "1) Открой Массив, кликни красную ячейку BP8 — в строке формул будет =Расчёт!O8",
        "2) Открой Расчёт, кликни O8 — увидишь урезку до лимита канала",
        "3) Кликни недельные столбцы справа на Расчёте — увидишь MIN(прогноз; остаток) и перенос",
        "4) Строка 4 на Массиве — СУММ плана по каналу; сравни с целями в строке 5 (тонны)",
        "",
        "ЛОГИКА",
        "Остаток + выпуск недели -> раздача по приоритету каналов -> остаток на следующую неделю.",
        "Сумма плана канала не больше цели (тонны*1000).",
        "План производства не менялся.",
        "",
        "ЦВЕТА",
        "Красные на Массиве — требуемый ответ «План мес».",
        "Зелёные на Расчёте и в строке 4 — формулы расчёта/проверки.",
    ]
    for i, t in enumerate(lines, 1):
        readme.cell(i, 1, t)
        readme.cell(i, 1).font = FONT_B if t.isupper() or t.startswith("ЧТО") or t.startswith("КАК") or t.startswith("ЛОГИКА") or t.startswith("ЦВЕТА") else FONT_N
        readme.cell(i, 1).fill = FILL_NOTE
    readme.column_dimensions["A"].width = 100

    print("Saving", OUT)
    wb.save(OUT)
    print("OK")


if __name__ == "__main__":
    main()
