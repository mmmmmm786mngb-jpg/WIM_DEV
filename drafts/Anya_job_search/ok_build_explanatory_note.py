#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Build separate explanatory note Excel for OK Demand test."""

import json
from pathlib import Path

import pythoncom
import win32com.client
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

SRC = Path(r"c:\Users\Acer\Downloads\OK_test_Demand") / "Тест_Деманд_ФОРМУЛЫ_Массив.xlsx"
OUT = Path(r"c:\Users\Acer\Downloads\OK_test_Demand") / "Poyasnitelnaya_zapiska_Demand.xlsx"
TMP = Path(r"c:\1c\Cursor_1c\WIM_DEV\drafts\Anya_job_search") / "ok_deficits.json"

FILL_Y = PatternFill("solid", fgColor="FFF2CC")
FILL_O = PatternFill("solid", fgColor="F4B183")
FILL_G = PatternFill("solid", fgColor="C6EFCE")
FILL_B = PatternFill("solid", fgColor="DDEBF7")
FILL_H = PatternFill("solid", fgColor="305496")
FILL_W = PatternFill("solid", fgColor="FCE4D6")
FONT_T = Font(bold=True, name="Calibri", size=14, color="1F4E79")
FONT_H = Font(bold=True, color="FFFFFF", name="Calibri", size=11)
FONT_B = Font(bold=True, name="Calibri", size=11)
FONT_N = Font(name="Calibri", size=11)
THIN = Border(
    left=Side(style="thin", color="B0B0B0"),
    right=Side(style="thin", color="B0B0B0"),
    top=Side(style="thin", color="B0B0B0"),
    bottom=Side(style="thin", color="B0B0B0"),
)


def extract_deficits():
    pythoncom.CoInitialize()
    xl = win32com.client.DispatchEx("Excel.Application")
    xl.Visible = False
    xl.DisplayAlerts = False
    wb = xl.Workbooks.Open(str(SRC))
    xl.CalculateFull()
    calc = wb.Worksheets("Расчёт")
    massiv = wb.Worksheets("Массив")
    rows = []
    for r in range(8, 436):
        st = calc.Range("E{0}".format(r)).Value
        if st != "ДЕФИЦИТ":
            continue
        code = calc.Range("A{0}".format(r)).Value
        supply = calc.Range("B{0}".format(r)).Value or 0
        demand = calc.Range("C{0}".format(r)).Value or 0
        gap = calc.Range("D{0}".format(r)).Value or 0
        plan = calc.Range("W{0}".format(r)).Value or 0
        cat = massiv.Range("D{0}".format(r)).Value
        atype = massiv.Range("E{0}".format(r)).Value
        fill = (plan / demand * 100.0) if demand else 100.0
        rows.append(
            {
                "row": r,
                "code": code,
                "cat": cat,
                "atype": atype,
                "supply": round(float(supply), 2),
                "demand": round(float(demand), 2),
                "gap": round(float(gap), 2),
                "gap_pct": round(float(gap) / float(demand) * 100.0, 2) if demand else 0,
                "plan": round(float(plan), 2),
                "fill": round(fill, 2),
            }
        )
    wb.Close(False)
    xl.Quit()
    pythoncom.CoUninitialize()
    rows.sort(key=lambda x: -x["gap"])
    TMP.write_text(json.dumps(rows, ensure_ascii=False, indent=2), encoding="utf-8")
    return rows


def write_note(ws, n_def, sum_gap):
    ws.sheet_view.showGridLines = False
    ws["A1"] = "Пояснительная записка к решению теста"
    ws["A1"].font = FONT_T
    ws["A1"].fill = FILL_Y
    ws.merge_cells("A1:B1")

    ws["A2"] = (
        "Вакансия: Специалист по прогнозированию спроса / ООО «Объединенные кондитеры». "
        "Файл решения: Тест_Деманд_ФОРМУЛЫ_Массив.xlsx (красные колонки «План мес» на листе Массив)."
    )
    ws["A2"].fill = FILL_Y
    ws["A2"].alignment = Alignment(wrap_text=True)
    ws.merge_cells("A2:B2")
    ws.row_dimensions[2].height = 40

    blocks = [
        (
            "1. Список дефицитных позиций",
            FILL_O,
            [
                "Дефицитными считаются SKU, где доступный остаток на начало периода + план производства "
                "по неделям меньше суммы неограниченного прогноза спроса по каналам.",
                "Полный список — на листе «01_Список_дефицита» (оранжевая заливка). "
                "Всего дефицитных позиций: {n}. Суммарный дефицит относительно прогноза: {gap:,.0f} кг "
                "({gap_t:,.1f} т).".format(n=n_def, gap=sum_gap, gap_t=sum_gap / 1000.0),
                "В списке указаны: код, категория, тип ассортимента, доступный объём, прогноз, "
                "величина дефицита, итоговый план после приоритизации и fill rate.",
            ],
        ),
        (
            "2. Причина распределения дефицитных позиций именно так",
            FILL_Y,
            [
                "При нехватке товара объём распределяется по приоритету каналов "
                "(от высшего к низшему):",
                "1) ФС промо — обязательства / промо федеральных сетей;",
                "2) ФС рег — регуляр федеральных сетей;",
                "3) РДЦ — региональные РЦ, сервис регионов;",
                "4) РКП — розничная сеть компании;",
                "5) ЭК — экспорт (контрактные обязательства);",
                "6) ФТ — фирменная торговля;",
                "7) КПи — корпоративные продажи;",
                "8) E-commerce — более гибкий канал.",
                "Расчёт выполнен с недельным шагом: остаток начала + выпуск недели → отгрузка "
                "по приоритету → переходящий остаток на следующую неделю. "
                "План производства и план выпуска не увеличивались и не снимались.",
                "После сборки «сырого» плана сумма по каждому каналу ограничивается целевым планом "
                "канала из шапки таблицы (тонны × 1000 = кг) пропорциональным коэффициентом, "
                "чтобы не выйти за рамки суммарного плана канала.",
                "Прогноз спроса рассматривался как неограниченное видение отдела продаж, не как заказ. "
                "Целевой остаток не использовался (согласно условию задания).",
            ],
        ),
        (
            "3. Риски при образовании дефицита",
            FILL_W,
            [
                "Недопоставка в федеральные сети: штрафы, cut-off промо, потеря полки и ranking.",
                "Просадка OTIF и уровня сервиса РДЦ: локальные out-of-stock, недовольство регионов.",
                "Искажение сигнала спроса: если резать каналы без явного приоритета, продажи усилят "
                "overforecasting в следующем цикле планирования.",
                "Перенос спроса на аналоги и конкурентов; риск потери доли рынка по ключевым SKU.",
                "Одновременно — заморозка оборотного капитала в профицитных остатках по другим позициям, "
                "если дефицит хитов не сопровождается работой с излишками.",
            ],
        ),
        (
            "4. Действия по позициям, которые могут находиться в дефиците",
            FILL_B,
            [
                "1) Зафиксировать short-list дефицита и согласовать с Sales / Marketing приоритет "
                "каналов на текущий S&OP-цикл (защита ФС промо и ключевых сетей в первую очередь).",
                "2) Запросить у Production возможность переброски мощностей / сырья на дефицитные SKU "
                "внутри утверждённого плана выпуска (без увеличения общего лимита, если это запрещено).",
                "3) По позициям с устойчивым дефицитом: сократить / не подтверждать extra-промо и "
                "доп. заказы; предложить субституты из профицитного ассортимента.",
                "4) Профицитные объёмы направлять в каналы с недобором до целевого плана канала "
                "(строго в пределах суммарного плана канала).",
                "5) Weekly-контроль: факт отгрузок vs план, early-warning по отрицательным "
                "прогнозным остаткам на 1–2 недели вперёд.",
                "6) Эскалация на S&OP с decision log: accept OOS / cut channel / replan production.",
            ],
        ),
        (
            "Краткая справка по файлу решения",
            FILL_G,
            [
                "Лист Массив, красные колонки BO:BX «План мес» — итоговый ответ, заполнены формулами.",
                "Лист Расчёт — промежуточные формулы (недельный приоритет, перенос остатка, урезка до лимита).",
                "Строка 4 листа Массив — сумма плана по каналу для сверки с целями в строке 5 (тонны).",
            ],
        ),
    ]

    r = 4
    for title, fill, paras in blocks:
        ws.cell(r, 1, title)
        ws.cell(r, 1).font = FONT_B
        ws.cell(r, 1).fill = fill
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
        r += 1
        for p in paras:
            ws.cell(r, 1, p)
            ws.cell(r, 1).font = FONT_N
            ws.cell(r, 1).fill = fill
            ws.cell(r, 1).alignment = Alignment(wrap_text=True, vertical="top")
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
            ws.row_dimensions[r].height = 36 if len(p) > 80 else 20
            r += 1
        r += 1

    ws.column_dimensions["A"].width = 100
    ws.column_dimensions["B"].width = 12


def write_deficit_sheet(ws, rows):
    ws["A1"] = "Приложение: список дефицитных позиций"
    ws["A1"].font = FONT_T
    ws["A1"].fill = FILL_O
    ws.merge_cells("A1:J1")

    ws["A2"] = (
        "Источник: расчёт по файлу Тест_Деманд_ФОРМУЛЫ_Массив.xlsx. "
        "Критерий: доступно (остаток+выпуск) < сумма прогноза каналов по неделям."
    )
    ws["A2"].fill = FILL_Y
    ws.merge_cells("A2:J2")

    headers = [
        "№",
        "Строка_Массив",
        "Короткий код",
        "Категория",
        "Тип ассортимента",
        "Доступно, кг",
        "Прогноз, кг",
        "Дефицит, кг",
        "Дефицит, %",
        "План после приоритета, кг",
        "Fill rate, %",
    ]
    for c, h in enumerate(headers, 1):
        cell = ws.cell(4, c, h)
        cell.fill = FILL_H
        cell.font = FONT_H
        cell.alignment = Alignment(wrap_text=True, horizontal="center")
        cell.border = THIN
    ws.row_dimensions[4].height = 32

    for i, d in enumerate(rows, 1):
        vals = [
            i,
            d["row"],
            d["code"],
            d["cat"],
            d["atype"],
            d["supply"],
            d["demand"],
            d["gap"],
            d["gap_pct"],
            d["plan"],
            d["fill"],
        ]
        for c, v in enumerate(vals, 1):
            cell = ws.cell(4 + i, c, v)
            cell.fill = FILL_O
            cell.border = THIN
            cell.font = FONT_N

    end = 4 + len(rows)
    ws.cell(end + 2, 1, "Итого позиций:")
    ws.cell(end + 2, 2, len(rows))
    ws.cell(end + 2, 1).fill = FILL_Y
    ws.cell(end + 2, 2).fill = FILL_Y
    ws.cell(end + 2, 1).font = FONT_B
    ws.cell(end + 2, 7, "Сумма дефицита, кг:")
    ws.cell(end + 2, 8, round(sum(d["gap"] for d in rows), 2))
    ws.cell(end + 2, 7).fill = FILL_Y
    ws.cell(end + 2, 8).fill = FILL_Y
    ws.cell(end + 2, 7).font = FONT_B

    ws.auto_filter.ref = "A4:K{0}".format(end)
    ws.freeze_panes = "C5"
    widths = [6, 12, 14, 14, 16, 14, 14, 12, 12, 18, 12]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def write_summary(ws, rows):
    ws["A1"] = "Краткая сводка для письма HR"
    ws["A1"].font = FONT_T
    ws["A1"].fill = FILL_Y

    text = (
        "Добрый день!\n\n"
        "Направляю решение тестового задания и пояснительную записку.\n"
        "В файле Тест_Деманд_ФОРМУЛЫ_Массив.xlsx заполнен блок «План мес» (красные колонки) "
        "с учётом недельного переноса остатков и приоритета каналов при дефиците. "
        "План производства не изменялся; суммарные планы каналов не превышают целевые значения из шапки.\n\n"
        "Дефицитных позиций: {n}. Подробный список и логика — в файле пояснительной записки.\n\n"
        "С уважением,\nАнна"
    ).format(n=len(rows))
    ws["A3"] = text
    ws["A3"].alignment = Alignment(wrap_text=True, vertical="top")
    ws["A3"].fill = FILL_B
    ws.merge_cells("A3:B12")
    ws.column_dimensions["A"].width = 90
    ws.row_dimensions[3].height = 160


def main():
    print("Extract deficits from calculated file...")
    rows = extract_deficits()
    print("deficits", len(rows))
    sum_gap = sum(d["gap"] for d in rows)

    wb = Workbook()
    ws0 = wb.active
    ws0.title = "00_Пояснительная"
    write_note(ws0, len(rows), sum_gap)

    ws1 = wb.create_sheet("01_Список_дефицита")
    write_deficit_sheet(ws1, rows)

    ws2 = wb.create_sheet("02_Черновик_письма")
    write_summary(ws2, rows)

    print("Saving", OUT)
    wb.save(OUT)
    print("OK")


if __name__ == "__main__":
    main()
