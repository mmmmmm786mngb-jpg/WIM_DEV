#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Copy Test_Demand_FORMULY_Massiv and add explanatory note sheets inside.
"""

from pathlib import Path
from shutil import copy2

import pythoncom
import win32com.client
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

SRC = Path(r"c:\Users\Acer\Downloads\OK_test_Demand") / "Тест_Деманд_ФОРМУЛЫ_Массив.xlsx"
OUT = Path(r"c:\Users\Acer\Downloads\OK_test_Demand") / "Тест_Деманд_ФОРМУЛЫ_Массив_с_запиской.xlsx"

FILL_Y = PatternFill("solid", fgColor="FFF2CC")
FILL_O = PatternFill("solid", fgColor="F4B183")
FILL_G = PatternFill("solid", fgColor="C6EFCE")
FILL_B = PatternFill("solid", fgColor="DDEBF7")
FILL_H = PatternFill("solid", fgColor="305496")
FILL_W = PatternFill("solid", fgColor="FCE4D6")
FONT_T = Font(bold=True, name="Calibri", size=14, color="1F4E79")
FONT_H = Font(bold=True, color="FFFFFF", name="Calibri", size=11)
FONT_B = Font(bold=True, name="Calibri", size=11)
FONT_N = Font(name="Calibri", size=11)
THIN = Border(
    left=Side(style="thin", color="B0B0B0"),
    right=Side(style="thin", color="B0B0B0"),
    top=Side(style="thin", color="B0B0B0"),
    bottom=Side(style="thin", color="B0B0B0"),
)


def extract_deficits(path):
    pythoncom.CoInitialize()
    xl = win32com.client.DispatchEx("Excel.Application")
    xl.Visible = False
    xl.DisplayAlerts = False
    wb = xl.Workbooks.Open(str(path))
    xl.CalculateFull()
    calc = wb.Worksheets("Расчёт")
    massiv = wb.Worksheets("Массив")
    rows = []
    for r in range(8, 436):
        if calc.Range("E{0}".format(r)).Value != "ДЕФИЦИТ":
            continue
        code = calc.Range("A{0}".format(r)).Value
        supply = float(calc.Range("B{0}".format(r)).Value or 0)
        demand = float(calc.Range("C{0}".format(r)).Value or 0)
        gap = float(calc.Range("D{0}".format(r)).Value or 0)
        plan = float(calc.Range("W{0}".format(r)).Value or 0)
        rows.append(
            {
                "row": r,
                "code": code,
                "cat": massiv.Range("D{0}".format(r)).Value,
                "atype": massiv.Range("E{0}".format(r)).Value,
                "supply": round(supply, 2),
                "demand": round(demand, 2),
                "gap": round(gap, 2),
                "gap_pct": round(gap / demand * 100.0, 2) if demand else 0.0,
                "plan": round(plan, 2),
                "fill": round(plan / demand * 100.0, 2) if demand else 100.0,
            }
        )
    wb.Close(False)
    xl.Quit()
    pythoncom.CoUninitialize()
    rows.sort(key=lambda x: -x["gap"])
    return rows


def add_note_sheet(wb, n_def, sum_gap):
    name = "Пояснительная_записка"
    if name in wb.sheetnames:
        del wb[name]
    ws = wb.create_sheet(name, 0)
    ws.sheet_view.showGridLines = False

    ws["A1"] = "Пояснительная записка к решению теста"
    ws["A1"].font = FONT_T
    ws["A1"].fill = FILL_Y
    ws.merge_cells("A1:B1")

    ws["A2"] = (
        "Вакансия: Специалист по прогнозированию спроса / ООО «Объединенные кондитеры». "
        "Решение: красные колонки «План мес» (BO:BX) на листе Массив — заполнены формулами; "
        "промежуточный расчёт — лист Расчёт."
    )
    ws["A2"].fill = FILL_Y
    ws["A2"].alignment = Alignment(wrap_text=True)
    ws.merge_cells("A2:B2")
    ws.row_dimensions[2].height = 45

    blocks = [
        (
            "1. Список дефицитных позиций",
            FILL_O,
            [
                "Дефицитными считаются SKU, где доступный остаток на начало периода + план производства "
                "по неделям меньше суммы неограниченного прогноза спроса по каналам.",
                "Полный список — на листе «Список_дефицита». Всего дефицитных позиций: {n}. "
                "Суммарный дефицит относительно прогноза: {gap:,.0f} кг ({gap_t:,.1f} т).".format(
                    n=n_def, gap=sum_gap, gap_t=sum_gap / 1000.0
                ),
                "В списке: код, категория, тип ассортимента, доступно, прогноз, дефицит, "
                "план после приоритизации, fill rate.",
            ],
        ),
        (
            "2. Причина распределения дефицитных позиций так, а не иначе",
            FILL_Y,
            [
                "При нехватке товара объём распределяется по приоритету каналов (от высшего к низшему):",
                "1) ФС промо — обязательства / промо федеральных сетей;",
                "2) ФС рег — регуляр федеральных сетей;",
                "3) РДЦ — региональные РЦ, сервис регионов;",
                "4) РКП — розничная сеть компании;",
                "5) ЭК — экспорт (контрактные обязательства);",
                "6) ФТ — фирменная торговля;",
                "7) КПи — корпоративные продажи;",
                "8) E-commerce — более гибкий канал.",
                "Расчёт с недельным шагом: остаток начала + выпуск недели -> отгрузка по приоритету -> "
                "переходящий остаток на следующую неделю. План производства и план выпуска не увеличивались.",
                "После сборки сырого плана сумма по каналу ограничивается целевым планом канала "
                "(тонны из строки 5 Массива * 1000 = кг), чтобы не выйти за суммарный план канала.",
                "Прогноз спроса — неограниченное видение продаж, не заказ. Целевой остаток не использовался.",
            ],
        ),
        (
            "3. Какие риски видит кандидат при образовании дефицита",
            FILL_W,
            [
                "Недопоставка в федеральные сети: штрафы, cut-off промо, потеря полки и ranking.",
                "Просадка OTIF и уровня сервиса РДЦ: локальные out-of-stock в регионах.",
                "Искажение сигнала спроса: резка без приоритета усиливает overforecasting в следующем цикле.",
                "Перенос спроса на аналоги и конкурентов; риск потери доли рынка по ключевым SKU.",
                "Заморозка оборотного капитала в профицитных остатках при одновременном дефиците хитов.",
            ],
        ),
        (
            "4. Какие действия будут предприняты по позициям в дефиците",
            FILL_B,
            [
                "1) Short-list дефицита на S&OP с Sales/Marketing; защита ФС промо и ключевых сетей.",
                "2) Запрос Production на переброску мощностей/сырья внутри утверждённого плана выпуска.",
                "3) Сокращение / неподтверждение extra-промо и доп. заказов по хроническому дефициту; "
                "субституты из профицитного ассортимента.",
                "4) Профицит направлять в каналы с недобором до целевого плана канала (в пределах лимита).",
                "5) Weekly-контроль остатка на 1-2 недели вперёд (early-warning).",
                "6) Эскалация на S&OP: accept OOS / cut channel / replan production — с decision log.",
            ],
        ),
    ]

    r = 4
    for title, fill, paras in blocks:
        ws.cell(r, 1, title).font = FONT_B
        ws.cell(r, 1).fill = fill
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
        r += 1
        for p in paras:
            cell = ws.cell(r, 1, p)
            cell.font = FONT_N
            cell.fill = fill
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
            ws.row_dimensions[r].height = 36 if len(p) > 90 else 20
            r += 1
        r += 1

    ws.column_dimensions["A"].width = 105
    ws.column_dimensions["B"].width = 12
    return ws


def add_deficit_sheet(wb, rows):
    name = "Список_дефицита"
    if name in wb.sheetnames:
        del wb[name]
    ws = wb.create_sheet(name, 1)

    ws["A1"] = "Приложение к пояснительной записке: список дефицитных позиций"
    ws["A1"].font = FONT_T
    ws["A1"].fill = FILL_O
    ws.merge_cells("A1:K1")

    ws["A2"] = (
        "Критерий: доступно (остаток + выпуск по неделям) < сумма прогноза каналов. "
        "Статус считается на листе Расчёт (столбец Статус)."
    )
    ws["A2"].fill = FILL_Y
    ws.merge_cells("A2:K2")

    headers = [
        "№",
        "Строка_Массив",
        "Короткий код",
        "Категория",
        "Тип ассортимента",
        "Доступно, кг",
        "Прогноз, кг",
        "Дефицит, кг",
        "Дефицит, %",
        "План после приоритета, кг",
        "Fill rate, %",
    ]
    for c, h in enumerate(headers, 1):
        cell = ws.cell(4, c, h)
        cell.fill = FILL_H
        cell.font = FONT_H
        cell.alignment = Alignment(wrap_text=True, horizontal="center")
        cell.border = THIN
    ws.row_dimensions[4].height = 32

    for i, d in enumerate(rows, 1):
        vals = [
            i,
            d["row"],
            d["code"],
            d["cat"],
            d["atype"],
            d["supply"],
            d["demand"],
            d["gap"],
            d["gap_pct"],
            d["plan"],
            d["fill"],
        ]
        for c, v in enumerate(vals, 1):
            cell = ws.cell(4 + i, c, v)
            cell.fill = FILL_O
            cell.border = THIN
            cell.font = FONT_N

    end = 4 + len(rows)
    ws.cell(end + 2, 1, "Итого позиций:").font = FONT_B
    ws.cell(end + 2, 2, len(rows))
    ws.cell(end + 2, 1).fill = FILL_Y
    ws.cell(end + 2, 2).fill = FILL_Y
    ws.cell(end + 2, 7, "Сумма дефицита, кг:").font = FONT_B
    ws.cell(end + 2, 8, round(sum(d["gap"] for d in rows), 2))
    ws.cell(end + 2, 7).fill = FILL_Y
    ws.cell(end + 2, 8).fill = FILL_Y

    ws.auto_filter.ref = "A4:K{0}".format(end)
    ws.freeze_panes = "C5"
    widths = [6, 12, 14, 14, 16, 14, 14, 12, 12, 18, 12]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def main():
    print("Extract deficits...")
    rows = extract_deficits(SRC)
    print("deficits", len(rows))
    sum_gap = sum(d["gap"] for d in rows)

    print("Copy file...")
    if OUT.exists():
        OUT.unlink()
    copy2(SRC, OUT)

    print("Open copy and add note sheets...")
    wb = load_workbook(OUT)
    add_note_sheet(wb, len(rows), sum_gap)
    add_deficit_sheet(wb, rows)

    # bump readme if exists
    if "00_Как_сдавать" in wb.sheetnames:
        ws = wb["00_Как_сдавать"]
        ws["A20"] = "Добавлено: листы «Пояснительная_записка» и «Список_дефицита»."
        ws["A20"].fill = FILL_G
        ws["A20"].font = FONT_B

    print("Saving", OUT)
    wb.save(OUT)
    print("OK")


if __name__ == "__main__":
    main()
